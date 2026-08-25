"""
drake_prefill_importer.py
-------------------------
Ingest Drake CSM exports + purple-sheet form chunks into prefill tables.

Phase 1: composite CSM dedupe + disposition flags.
Phase 2: purple parse / JSON counts / name match (dry-run before prod writes).

Fuzzy accept threshold is local (90) — do not change name_matcher.ACCEPT_THRESHOLD.
"""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import logging
import re
import sqlite3
import sys
from collections import Counter, defaultdict
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path
from typing import Any, Iterable, Optional

from rapidfuzz import fuzz

from utils import now as utc_now

log = logging.getLogger(__name__)

# Local to this importer — do not change name_matcher.ACCEPT_THRESHOLD (88).
PREFILL_FUZZY_ACCEPT = 90

# Lower rank = more specific claim. primary_only must lose every tie.
VARIANT_SPECIFICITY_RANK: dict[str, int] = {
    "raw_normalized": 0,
    "raw_no_comma": 0,
    "shared_surname_joint": 1,
    "separate_surname_joint": 1,
    "shared_surname_joint_mi_stripped": 2,
    "separate_surname_joint_mi_stripped": 2,
    "single_flip": 3,
    "fallback_flip_part": 3,
    "single_flip_mi_stripped": 4,
    "primary_only_flip": 5,
    "primary_only_flip_mi_stripped": 5,
}

_EPOCH_MIN = datetime(1, 1, 1)

_DISPOSITION_BY_STATUS: dict[str, str] = {
    "EF ACCEPTED": "PY_FILED_ACCEPTED",
    "EF REJECTED": "PY_REJECTED",
    "EF EXT ACCEPTED": "PY_EXTENDED",
    "IN PROGRESS": "PY_INCOMPLETE",
    "PRINTED": "PY_INCOMPLETE",
    "UPDATED FROM 2023": "PY_ROLLOVER_ONLY",
}

_BOOL_INDICATOR_COLS = frozenset({"Form 4868 Indicator", "Form 2350 Indicator"})
_DECIMAL_COLS = frozenset({"Schedule 1A Deduction"})
_TEXT_COLS = frozenset({"Return Type", "Taxpayer Name"})

_EXPECTED_CHUNK_COLS = (23, 24, 25)
_EXPECTED_DATA_ROWS = 1298


# ── Phase 1 dataclasses ──────────────────────────────────────────────────────


@dataclass
class CsmRow:
    ssn_last4: str
    name_raw: str
    name_norm: str
    status: Optional[str]
    last_change: datetime
    last_change_raw: Optional[str]
    started: Optional[str] = None
    completed: Optional[str] = None
    return_type: Optional[str] = None
    changed_by: Optional[str] = None
    amount_owed: Optional[str] = None
    raw: dict[str, Any] = field(default_factory=dict)


@dataclass
class StatusRecovery:
    ssn_last4: str
    name_raw: str
    recovered_status: str
    status_as_of: Optional[str]
    anchor_changed: Optional[str]


@dataclass
class CsmDedupeResult:
    rows_read: int
    clients_after: int
    clients: list[dict[str, Any]]
    status_counts: Counter
    disposition_counts: Counter
    recoveries: list[StatusRecovery]
    recovery_status_counts: Counter
    history_groups: int
    empty_name_or_ssn_skipped: int


@dataclass
class PurpleJoined:
    taxpayer_name: str
    name_norm: str
    form_counts: dict[str, Any]
    return_type: Optional[str]
    source_files: list[str]
    row_index: int
    ambiguous_duplicate_name: bool = False


@dataclass
class PrefillLinkDraft:
    csm: dict[str, Any]
    prefill_status: str
    match_tier: Optional[str] = None
    match_score: Optional[float] = None
    matched_variant: Optional[str] = None
    purple_name: Optional[str] = None
    form_counts: Optional[dict[str, Any]] = None
    return_type: Optional[str] = None
    source_files: Optional[list[str]] = None
    reason: Optional[str] = None


@dataclass
class CollisionStats:
    """Deterministic purple-name collision accounting (unique purple names only)."""

    uncontested: int = 0
    contested: int = 0
    resolved_by_precedence: int = 0
    still_tied: int = 0
    no_csm_hit: int = 0
    fuzzy_multi_claim_groups: int = 0
    fuzzy_resolved_by_precedence: int = 0
    fuzzy_resolved_by_score: int = 0
    fuzzy_still_tied: int = 0


@dataclass
class Phase2DryRun:
    csm: CsmDedupeResult
    chunk_row_counts: list[int]
    chunk_col_counts: list[int]
    purple_unique_names: int
    purple_duplicate_names: list[tuple[str, int]]
    links: list[PrefillLinkDraft]
    tier_counts: Counter
    variant_counts: Counter
    prefill_status_counts: Counter
    no_form_by_disposition: Counter
    collision_stats: CollisionStats
    purple_return_type_counts: Counter
    matched_return_type_counts: Counter
    assertions_ok: list[str]


@dataclass
class PrefillWriteStats:
    inserted: int = 0
    updated: int = 0
    skipped_manual: int = 0
    form_rows_upserted: int = 0
    form_rows_deleted: int = 0
    batch_id: Optional[int] = None


def variant_rank(label: str) -> int:
    return VARIANT_SPECIFICITY_RANK.get(label, 99)


# ── Shared helpers ────────────────────────────────────────────────────────────


def normalize_csm_name(raw: str) -> str:
    """Uppercase + collapse whitespace (identity key component).

    Also normalizes comma spacing so CSM history twins like
    ``MARTINEZ , FERNANDO`` and ``MARTINEZ, FERNANDO`` collapse to one key.
    Without this, both claim the same purple joint name and land in manual —
    including many genuine 1040SR returns.
    """
    s = re.sub(r"\s+", " ", (raw or "").upper().strip())
    s = re.sub(r"\s*,\s*", ", ", s)
    return s


def normalize_purple_name(raw: str) -> str:
    return normalize_csm_name(raw)


def extract_ssn_last4(id_value: Any) -> Optional[str]:
    if id_value is None:
        return None
    digits = "".join(c for c in str(id_value) if c.isdigit())
    return digits[-4:] if len(digits) >= 4 else None


def parse_last_change(value: Any) -> tuple[datetime, Optional[str]]:
    if value is None or value == "":
        return _EPOCH_MIN, None
    if isinstance(value, datetime):
        return value.replace(tzinfo=None), value.strftime("%m/%d/%Y %H:%M:%S")
    raw = str(value).strip()
    if not raw:
        return _EPOCH_MIN, None
    for fmt in ("%m/%d/%Y %H:%M:%S", "%m/%d/%Y"):
        try:
            return datetime.strptime(raw, fmt), raw
        except ValueError:
            continue
    return _EPOCH_MIN, raw


def disposition_for_status(status: Optional[str]) -> str:
    if not status or not str(status).strip():
        return "PY_STATUS_UNKNOWN"
    key = re.sub(r"\s+", " ", str(status).strip().upper())
    return _DISPOSITION_BY_STATUS.get(key, "PY_STATUS_UNKNOWN")


def _cell_str(value: Any) -> Optional[str]:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime("%m/%d/%Y %H:%M:%S")
    s = str(value).strip()
    return s or None


# ── Phase 1: CSM ──────────────────────────────────────────────────────────────


def load_csm_xlsx(path: str | Path) -> list[CsmRow]:
    # openpyxl is imported here, not at module scope: app.py imports helpers from
    # this module during a request, and the service host has no openpyxl.
    from openpyxl import load_workbook

    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"CSM file not found: {path}")

    wb = load_workbook(path, read_only=True, data_only=True)
    try:
        ws = wb.active
        rows_iter = ws.iter_rows(values_only=True)
        header_row = next(rows_iter, None)
        if not header_row:
            raise ValueError("CSM workbook is empty")
        headers = [(_cell_str(h) or f"COL_{i}") for i, h in enumerate(header_row)]
        header_upper = {h.upper(): h for h in headers}
        required = {"ID (LAST 4)", "CLIENT NAME", "STATUS", "LAST CHANGE"}
        missing = required - set(header_upper)
        if missing:
            raise ValueError(
                f"CSM header missing columns {sorted(missing)}; got {headers}"
            )

        out: list[CsmRow] = []
        for vals in rows_iter:
            if vals is None or all(v is None or str(v).strip() == "" for v in vals):
                continue
            row = {
                headers[i]: vals[i] if i < len(vals) else None
                for i in range(len(headers))
            }
            name_raw = _cell_str(row.get(header_upper["CLIENT NAME"])) or ""
            ssn = extract_ssn_last4(row.get(header_upper["ID (LAST 4)"]))
            status = _cell_str(row.get(header_upper["STATUS"]))
            lc_dt, lc_raw = parse_last_change(row.get(header_upper["LAST CHANGE"]))
            out.append(
                CsmRow(
                    ssn_last4=ssn or "",
                    name_raw=name_raw,
                    name_norm=normalize_csm_name(name_raw),
                    status=status,
                    last_change=lc_dt,
                    last_change_raw=lc_raw,
                    started=_cell_str(row.get(header_upper.get("STARTED", "Started"))),
                    completed=_cell_str(
                        row.get(header_upper.get("COMPLETED", "Completed"))
                    ),
                    return_type=_cell_str(row.get(header_upper.get("TYPE", "Type"))),
                    changed_by=_cell_str(
                        row.get(header_upper.get("CHANGED BY", "Changed By"))
                    ),
                    amount_owed=_cell_str(
                        row.get(header_upper.get("AMOUNT OWED", "Amount Owed"))
                    ),
                    raw={k: _cell_str(v) for k, v in row.items()},
                )
            )
        return out
    finally:
        wb.close()


def composite_dedupe_csm(rows: Iterable[CsmRow]) -> CsmDedupeResult:
    groups: dict[tuple[str, str], list[CsmRow]] = defaultdict(list)
    skipped = 0
    rows_list = list(rows)
    for r in rows_list:
        if not r.ssn_last4 or not r.name_norm:
            skipped += 1
            continue
        groups[(r.ssn_last4, r.name_norm)].append(r)

    clients: list[dict[str, Any]] = []
    recoveries: list[StatusRecovery] = []
    status_counts: Counter = Counter()
    disposition_counts: Counter = Counter()
    recovery_status_counts: Counter = Counter()
    history_groups = 0

    for (ssn, name_norm), group in groups.items():
        if len(group) > 1:
            history_groups += 1
        group_sorted = sorted(
            group,
            key=lambda x: (x.last_change, x.last_change_raw or ""),
            reverse=True,
        )
        anchor = group_sorted[0]
        status_row = next((g for g in group_sorted if g.status), None)
        resolved_status = status_row.status if status_row else None
        status_as_of = status_row.last_change_raw if status_row else None

        recovered = bool(
            status_row
            and (not anchor.status or not str(anchor.status).strip())
            and resolved_status
        )
        if recovered:
            recoveries.append(
                StatusRecovery(
                    ssn_last4=ssn,
                    name_raw=anchor.name_raw,
                    recovered_status=resolved_status or "",
                    status_as_of=status_as_of,
                    anchor_changed=anchor.last_change_raw,
                )
            )
            recovery_status_counts[resolved_status or "(empty)"] += 1

        disposition = disposition_for_status(resolved_status)
        status_key = (resolved_status or "").strip() or "(no status ever)"
        status_counts[status_key] += 1
        disposition_counts[disposition] += 1

        clients.append(
            {
                "csm_ssn_last4": ssn,
                "csm_name_raw": anchor.name_raw,
                "csm_name_norm": name_norm,
                "csm_status_raw": resolved_status,
                "csm_status_as_of": status_as_of,
                "csm_anchor_changed": anchor.last_change_raw,
                "disposition_status": disposition,
                "status_recovered_from_earlier_row": recovered,
                "history_row_count": len(group),
                "return_type": anchor.return_type,
                "started": anchor.started,
                "completed": anchor.completed,
                "changed_by": anchor.changed_by,
                "amount_owed": anchor.amount_owed,
            }
        )

    return CsmDedupeResult(
        rows_read=len(rows_list),
        clients_after=len(clients),
        clients=clients,
        status_counts=status_counts,
        disposition_counts=disposition_counts,
        recoveries=recoveries,
        recovery_status_counts=recovery_status_counts,
        history_groups=history_groups,
        empty_name_or_ssn_skipped=skipped,
    )


def run_csm_dry_run(csm_path: str | Path) -> CsmDedupeResult:
    rows = load_csm_xlsx(csm_path)
    result = composite_dedupe_csm(rows)
    if result.rows_read >= 2000 and result.history_groups < 10:
        raise RuntimeError(
            f"CSM dedupe removed almost no history "
            f"(rows={result.rows_read}, history_groups={result.history_groups})."
        )
    if result.rows_read >= 2000 and result.clients_after == result.rows_read:
        raise RuntimeError(
            "CSM dedupe kept 1:1 row count on a large file — expected history collapse."
        )
    return result


# ── Phase 2: purple parse ─────────────────────────────────────────────────────


def _coerce_form_value(header: str, raw: Optional[str]) -> Any:
    """Return JSON-serializable value; None means empty cell (not applicable)."""
    if raw is None:
        return None
    s = str(raw).strip()
    if s == "":
        return None

    if header in _BOOL_INDICATOR_COLS:
        return True if s.upper() == "X" else True if s.upper() in ("1", "Y", "YES") else None

    if header in _DECIMAL_COLS:
        try:
            return float(Decimal(s.replace(",", "").replace("$", "")))
        except (InvalidOperation, ValueError) as exc:
            raise ValueError(f"Bad decimal in {header!r}: {raw!r}") from exc

    if header in _TEXT_COLS:
        return s

    # Integer form counts — preserve 0 vs empty (empty already returned None)
    try:
        return int(s)
    except ValueError as exc:
        raise ValueError(f"Bad integer count in {header!r}: {raw!r}") from exc


def load_purple_chunk(path: str | Path, *, expected_cols: int) -> tuple[list[str], list[list[str]]]:
    """
    Return (headers, data_rows_padded).
    Skips title + as-of rows; asserts header width and ragged-row rules.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"Purple chunk not found: {path}")

    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))

    if len(rows) < 3:
        raise ValueError(f"{path.name}: need title, as-of, header — got {len(rows)} lines")

    headers = rows[2]
    if len(headers) != expected_cols:
        raise ValueError(
            f"{path.name}: expected {expected_cols} header columns, got {len(headers)}: {headers}"
        )
    if not headers or headers[0].strip() != "Taxpayer Name":
        raise ValueError(
            f"{path.name}: header row 3 must start with Taxpayer Name, got {headers[:3]!r}"
        )

    data: list[list[str]] = []
    for i, row in enumerate(rows[3:], start=4):
        if not any((c or "").strip() for c in row):
            continue  # trailing blank
        if len(row) > len(headers):
            raise ValueError(
                f"{path.name} line {i}: row width {len(row)} > header {len(headers)}"
            )
        padded = list(row) + [""] * (len(headers) - len(row))
        data.append(padded)

    if len(data) != _EXPECTED_DATA_ROWS:
        raise ValueError(
            f"{path.name}: expected {_EXPECTED_DATA_ROWS} data rows, got {len(data)}"
        )
    return headers, data


def join_purple_chunks(
    chunk_paths: list[str | Path],
) -> tuple[list[PurpleJoined], list[str], list[int], list[int]]:
    """
    Load three chunks, assert equal name sequences, join on Taxpayer Name.
    Returns (joined_rows, assertion_notes, row_counts, col_counts).
    """
    if len(chunk_paths) != 3:
        raise ValueError("Expected exactly three purple chunk paths")

    loaded: list[tuple[str, list[str], list[list[str]]]] = []
    notes: list[str] = []
    row_counts: list[int] = []
    col_counts: list[int] = []

    for path, exp_cols in zip(chunk_paths, _EXPECTED_CHUNK_COLS):
        headers, data = load_purple_chunk(path, expected_cols=exp_cols)
        loaded.append((str(path), headers, data))
        row_counts.append(len(data))
        col_counts.append(len(headers))
        notes.append(f"{Path(path).name}: {len(data)} rows x {len(headers)} cols")

    if len(set(row_counts)) != 1:
        raise ValueError(f"Purple chunk row counts differ: {row_counts}")

    name_seqs = [[r[0].strip() for r in data] for _, _, data in loaded]
    if not (name_seqs[0] == name_seqs[1] == name_seqs[2]):
        # Find first divergence for a useful error
        for i, (a, b, c) in enumerate(zip(name_seqs[0], name_seqs[1], name_seqs[2])):
            if not (a == b == c):
                raise ValueError(
                    f"Purple Taxpayer Name sequences diverge at row {i}: "
                    f"{a!r} / {b!r} / {c!r}"
                )
        raise ValueError("Purple Taxpayer Name sequences diverge (length mismatch)")

    notes.append("Taxpayer Name sequences identical across all three chunks")

    name_counts = Counter(name_seqs[0])
    dup_names = {n for n, k in name_counts.items() if k > 1}

    joined: list[PurpleJoined] = []
    for idx in range(row_counts[0]):
        form_counts: dict[str, Any] = {}
        return_type: Optional[str] = None
        sources: list[str] = []
        taxpayer = name_seqs[0][idx]

        for path, headers, data in loaded:
            sources.append(Path(path).name)
            row = data[idx]
            for h, cell in zip(headers, row):
                if h == "Taxpayer Name":
                    continue
                # Key present with coerced value (null if empty)
                form_counts[h] = _coerce_form_value(h, cell if cell != "" else None)
                if h == "Return Type":
                    return_type = form_counts[h] if isinstance(form_counts[h], str) else None

        joined.append(
            PurpleJoined(
                taxpayer_name=taxpayer,
                name_norm=normalize_purple_name(taxpayer),
                form_counts=form_counts,
                return_type=return_type,
                source_files=sources,
                row_index=idx,
                ambiguous_duplicate_name=taxpayer in dup_names
                or normalize_purple_name(taxpayer)
                in {normalize_purple_name(d) for d in dup_names},
            )
        )

    # Mark all rows whose name_norm is duplicated
    norm_counts = Counter(j.name_norm for j in joined)
    for j in joined:
        if norm_counts[j.name_norm] > 1:
            j.ambiguous_duplicate_name = True

    return joined, notes, row_counts, col_counts


# ── Phase 2: name variants + match ────────────────────────────────────────────


def _strip_single_letter_tokens(text: str) -> str:
    """Drop single-letter alphabetic tokens (middle initials)."""
    parts = [t for t in text.split() if not (len(t) == 1 and t.isalpha())]
    return " ".join(parts)


def csm_name_variants(csm_name_raw: str) -> list[tuple[str, str]]:
    """
    Return (variant_label, normalized_variant) pairs.

    If the same normalized string is produced by more than one transform for
    this CSM client, keep the lowest (most specific) specificity rank.
    Labels are stable for matched_variant auditing.
    """
    raw = (csm_name_raw or "").strip()
    # norm -> (label, rank)
    best: dict[str, tuple[str, int]] = {}

    def add(label: str, value: str) -> None:
        norm = normalize_purple_name(value)
        if not norm:
            return
        rank = variant_rank(label)
        prev = best.get(norm)
        if prev is None or rank < prev[1]:
            best[norm] = (label, rank)

    add("raw_normalized", raw)

    if "," not in raw:
        # Business / already FIRST LAST style
        add("raw_no_comma", raw)
        return [(label, norm) for norm, (label, _) in best.items()]

    # May be "LAST, FIRST" or "LAST1, FIRST1 & LAST2, FIRST2"
    ampersand_parts = re.split(r"\s*&\s*", raw)
    if len(ampersand_parts) == 1:
        last, _, first = ampersand_parts[0].partition(",")
        last, first = last.strip(), first.strip()
        if last and first:
            flipped = f"{first} {last}"
            add("single_flip", flipped)
            add("single_flip_mi_stripped", _strip_single_letter_tokens(flipped))
            # Same string as single_flip — rank keeps single_flip over primary_only
            add("primary_only_flip", flipped)
            add(
                "primary_only_flip_mi_stripped",
                _strip_single_letter_tokens(flipped),
            )
        return [(label, norm) for norm, (label, _) in best.items()]

    # Joint return
    left = ampersand_parts[0].strip()
    right = ampersand_parts[1].strip()
    # Primary-only from left side (weakest — loses ties to full-name claimants)
    if "," in left:
        plast, _, pfirst = left.partition(",")
        plast, pfirst = plast.strip(), pfirst.strip()
        if plast and pfirst:
            primary = f"{pfirst} {plast}"
            add("primary_only_flip", primary)
            add("primary_only_flip_mi_stripped", _strip_single_letter_tokens(primary))

    if "," in left and "," in right:
        # Separate-surname joint: LAST1, FIRST1 & LAST2, FIRST2
        l1, _, f1 = left.partition(",")
        l2, _, f2 = right.partition(",")
        l1, f1, l2, f2 = l1.strip(), f1.strip(), l2.strip(), f2.strip()
        if l1 and f1 and l2 and f2:
            sep = f"{f1} {l1} & {f2} {l2}"
            add("separate_surname_joint", sep)
            sep_mi = (
                f"{_strip_single_letter_tokens(f1)} {l1} & "
                f"{_strip_single_letter_tokens(f2)} {l2}"
            )
            add("separate_surname_joint_mi_stripped", sep_mi)
    elif "," in left and "," not in right:
        # Shared-surname joint: LAST, FIRST1 & FIRST2
        last, _, firsts = left.partition(",")
        last = last.strip()
        first1 = firsts.strip()
        first2 = right.strip()
        if last and first1 and first2:
            shared = f"{first1} & {first2} {last}"
            add("shared_surname_joint", shared)
            shared_mi = (
                f"{_strip_single_letter_tokens(first1)} & "
                f"{_strip_single_letter_tokens(first2)} {last}"
            )
            add("shared_surname_joint_mi_stripped", shared_mi)
    else:
        # Fallback: flip each comma side independently if present
        for part in ampersand_parts:
            if "," in part:
                last, _, first = part.partition(",")
                last, first = last.strip(), first.strip()
                if last and first:
                    add("fallback_flip_part", f"{first} {last}")

    return [(label, norm) for norm, (label, _) in best.items()]


def match_csm_to_purple(
    csm_clients: list[dict[str, Any]],
    purple_rows: list[PurpleJoined],
) -> tuple[list[PrefillLinkDraft], Counter, Counter, CollisionStats]:
    """Return (links, tier_counts, variant_counts, collision_stats)."""
    by_norm: dict[str, list[PurpleJoined]] = defaultdict(list)
    for p in purple_rows:
        by_norm[p.name_norm].append(p)

    ambiguous_purple_norms = {n for n, rows in by_norm.items() if len(rows) > 1}
    unique_purple = {n: rows[0] for n, rows in by_norm.items() if len(rows) == 1}

    csm_by_key: dict[tuple[str, str], dict[str, Any]] = {
        (c["csm_ssn_last4"], c["csm_name_norm"]): c for c in csm_clients
    }
    # variant_norm -> list of (csm_key, label, rank)
    variant_index: dict[str, list[tuple[tuple[str, str], str, int]]] = defaultdict(list)
    csm_variants: dict[tuple[str, str], list[tuple[str, str]]] = {}

    for key, c in csm_by_key.items():
        variants = csm_name_variants(c["csm_name_raw"])
        csm_variants[key] = variants
        for label, vnorm in variants:
            variant_index[vnorm].append((key, label, variant_rank(label)))

    matched_csm: set[tuple[str, str]] = set()
    matched_purple: set[str] = set()
    # Purple names contested at best rank (genuine ties) — unavailable for fuzzy
    tied_purple: set[str] = set()
    links_by_csm: dict[tuple[str, str], PrefillLinkDraft] = {}
    tier_counts: Counter = Counter()
    variant_counts: Counter = Counter()
    stats = CollisionStats()

    def _assign_deterministic(
        ck: tuple[str, str], label: str, prow: PurpleJoined, pnorm: str
    ) -> bool:
        if ck in matched_csm:
            return False
        links_by_csm[ck] = PrefillLinkDraft(
            csm=csm_by_key[ck],
            prefill_status="PRIOR_YEAR_FORMS_AVAILABLE",
            match_tier="deterministic",
            match_score=100.0,
            matched_variant=label,
            purple_name=prow.taxpayer_name,
            form_counts=prow.form_counts,
            return_type=prow.return_type,
            source_files=prow.source_files,
        )
        matched_csm.add(ck)
        matched_purple.add(pnorm)
        tier_counts["deterministic"] += 1
        variant_counts[label] += 1
        return True

    # ── Deterministic (with specificity precedence on collisions) ─────────
    for pnorm, prow in unique_purple.items():
        hits = variant_index.get(pnorm, [])
        # Per CSM: keep the lowest-rank label that hits this purple string
        best_by_csm: dict[tuple[str, str], tuple[str, int]] = {}
        for ck, label, rank in hits:
            prev = best_by_csm.get(ck)
            if prev is None or rank < prev[1]:
                best_by_csm[ck] = (label, rank)

        if not best_by_csm:
            stats.no_csm_hit += 1
            continue

        if len(best_by_csm) == 1:
            stats.uncontested += 1
            ck, (label, _) = next(iter(best_by_csm.items()))
            _assign_deterministic(ck, label, prow, pnorm)
            continue

        stats.contested += 1
        best_rank = min(r for _, r in best_by_csm.values())
        survivors = [
            (ck, lab) for ck, (lab, r) in best_by_csm.items() if r == best_rank
        ]
        if len(survivors) == 1:
            stats.resolved_by_precedence += 1
            ck, label = survivors[0]
            _assign_deterministic(ck, label, prow, pnorm)
            # Weaker-rank claimants stay unmatched (may still match their own purple)
            continue

        # Genuine ambiguity at the best specificity rank
        stats.still_tied += 1
        tied_purple.add(pnorm)
        for ck, label in survivors:
            if ck in matched_csm:
                continue
            links_by_csm[ck] = PrefillLinkDraft(
                csm=csm_by_key[ck],
                prefill_status="NEEDS_MANUAL_LINK",
                match_tier="manual",
                matched_variant=label,
                purple_name=prow.taxpayer_name,
                reason="deterministic_multi_csm_hit",
            )
            matched_csm.add(ck)
            tier_counts["manual"] += 1

    # CSM variants that hit ambiguous purple duplicate names
    for pnorm in ambiguous_purple_norms:
        for ck, label, _rank in variant_index.get(pnorm, []):
            if ck in matched_csm:
                continue
            links_by_csm[ck] = PrefillLinkDraft(
                csm=csm_by_key[ck],
                prefill_status="NEEDS_MANUAL_LINK",
                match_tier="manual",
                matched_variant=label,
                purple_name=by_norm[pnorm][0].taxpayer_name,
                reason="purple_duplicate_name",
            )
            matched_csm.add(ck)
            tier_counts["manual"] += 1

    # ── Fuzzy (≥90) among remaining ───────────────────────────────────────
    remaining_csm = [ck for ck in csm_by_key if ck not in matched_csm]
    remaining_purple = {
        n: p
        for n, p in unique_purple.items()
        if n not in matched_purple and n not in tied_purple
    }

    # purple_norm -> list of (csm_key, score, label, rank)
    claims: dict[str, list[tuple[tuple[str, str], float, str, int]]] = defaultdict(list)

    for ck in remaining_csm:
        variants = csm_variants[ck]
        best_p = None
        best_score = -1.0
        best_label = ""
        best_r = 99
        for pnorm in remaining_purple:
            for label, vnorm in variants:
                score = float(fuzz.token_sort_ratio(vnorm, pnorm))
                rank = variant_rank(label)
                # Prefer higher score; on equal score prefer more-specific (lower) rank
                if score > best_score or (score == best_score and rank < best_r):
                    best_score = score
                    best_p = pnorm
                    best_label = label
                    best_r = rank
        if best_p is not None and best_score >= PREFILL_FUZZY_ACCEPT:
            claims[best_p].append((ck, best_score, best_label, best_r))

    for pnorm, claimants in claims.items():
        if len(claimants) == 1:
            ck, score, label, _r = claimants[0]
            if ck in matched_csm:
                continue
            prow = remaining_purple[pnorm]
            links_by_csm[ck] = PrefillLinkDraft(
                csm=csm_by_key[ck],
                prefill_status="PRIOR_YEAR_FORMS_AVAILABLE",
                match_tier="fuzzy",
                match_score=score,
                matched_variant=label,
                purple_name=prow.taxpayer_name,
                form_counts=prow.form_counts,
                return_type=prow.return_type,
                source_files=prow.source_files,
            )
            matched_csm.add(ck)
            matched_purple.add(pnorm)
            tier_counts["fuzzy"] += 1
            variant_counts[label] += 1
            continue

        stats.fuzzy_multi_claim_groups += 1
        best_rank = min(r for *_, r in claimants)
        at_rank = [c for c in claimants if c[3] == best_rank]
        if len(at_rank) == 1:
            stats.fuzzy_resolved_by_precedence += 1
            ck, score, label, _r = at_rank[0]
            if ck in matched_csm:
                continue
            prow = remaining_purple[pnorm]
            links_by_csm[ck] = PrefillLinkDraft(
                csm=csm_by_key[ck],
                prefill_status="PRIOR_YEAR_FORMS_AVAILABLE",
                match_tier="fuzzy",
                match_score=score,
                matched_variant=label,
                purple_name=prow.taxpayer_name,
                form_counts=prow.form_counts,
                return_type=prow.return_type,
                source_files=prow.source_files,
            )
            matched_csm.add(ck)
            matched_purple.add(pnorm)
            tier_counts["fuzzy"] += 1
            variant_counts[label] += 1
            continue

        best_score = max(s for _, s, _, _ in at_rank)
        top = [c for c in at_rank if c[1] == best_score]
        if len(top) == 1:
            stats.fuzzy_resolved_by_score += 1
            ck, score, label, _r = top[0]
            if ck in matched_csm:
                continue
            prow = remaining_purple[pnorm]
            links_by_csm[ck] = PrefillLinkDraft(
                csm=csm_by_key[ck],
                prefill_status="PRIOR_YEAR_FORMS_AVAILABLE",
                match_tier="fuzzy",
                match_score=score,
                matched_variant=label,
                purple_name=prow.taxpayer_name,
                form_counts=prow.form_counts,
                return_type=prow.return_type,
                source_files=prow.source_files,
            )
            matched_csm.add(ck)
            matched_purple.add(pnorm)
            tier_counts["fuzzy"] += 1
            variant_counts[label] += 1
            continue

        # Still tied after rank + score — high-priority manual
        stats.fuzzy_still_tied += 1
        tied_purple.add(pnorm)
        for ck, score, label, _r in top:
            if ck in matched_csm:
                continue
            links_by_csm[ck] = PrefillLinkDraft(
                csm=csm_by_key[ck],
                prefill_status="NEEDS_MANUAL_LINK",
                match_tier="manual",
                match_score=score,
                matched_variant=label,
                purple_name=remaining_purple[pnorm].taxpayer_name,
                reason="fuzzy_multi_csm_claim",
            )
            matched_csm.add(ck)
            tier_counts["manual"] += 1

    # Refresh remaining purple after fuzzy assignments / ties
    remaining_purple = {
        n: p
        for n, p in unique_purple.items()
        if n not in matched_purple and n not in tied_purple
    }

    # ── Leftover CSM classification ───────────────────────────────────────
    # Match failures never land in NO_PRIOR_FORM_DATA.
    # Contested / multi-claim → already NEEDS_MANUAL_LINK.
    # Orphan purple rows still need a review candidate → LOW_CONFIDENCE_NO_MATCH
    # (greedy best-CSM-per-orphan, score < accept). Excess unmatched CSM past
    # orphan inventory → NO_PRIOR_FORM_DATA (genuine absence; floor ≈ 106).
    unmatched = [ck for ck in csm_by_key if ck not in matched_csm]
    low_conf_claimed: set[tuple[str, str]] = set()

    # Pair each orphan purple to its best remaining CSM (orphan count is small).
    for pnorm, prow in remaining_purple.items():
        best_ck: Optional[tuple[str, str]] = None
        best_score = -1.0
        best_label = ""
        best_r = 99
        for ck in unmatched:
            if ck in low_conf_claimed or ck in matched_csm:
                continue
            for label, vnorm in csm_variants[ck]:
                score = float(fuzz.token_sort_ratio(vnorm, pnorm))
                rank = variant_rank(label)
                if score > best_score or (score == best_score and rank < best_r):
                    best_score = score
                    best_ck = ck
                    best_label = label
                    best_r = rank
        if best_ck is None:
            continue
        if best_score >= PREFILL_FUZZY_ACCEPT:
            # Mutual-best edge case the first fuzzy pass missed — auto-accept
            links_by_csm[best_ck] = PrefillLinkDraft(
                csm=csm_by_key[best_ck],
                prefill_status="PRIOR_YEAR_FORMS_AVAILABLE",
                match_tier="fuzzy",
                match_score=best_score,
                matched_variant=best_label,
                purple_name=prow.taxpayer_name,
                form_counts=prow.form_counts,
                return_type=prow.return_type,
                source_files=prow.source_files,
            )
            matched_csm.add(best_ck)
            matched_purple.add(pnorm)
            low_conf_claimed.add(best_ck)
            tier_counts["fuzzy"] += 1
            variant_counts[best_label] += 1
            continue
        # Below-accept candidate for an unmatched purple — visible, low priority
        links_by_csm[best_ck] = PrefillLinkDraft(
            csm=csm_by_key[best_ck],
            prefill_status="LOW_CONFIDENCE_NO_MATCH",
            match_tier="manual",
            match_score=best_score,
            matched_variant=best_label,
            purple_name=prow.taxpayer_name,
            reason="fuzzy_below_accept",
        )
        low_conf_claimed.add(best_ck)
        matched_csm.add(best_ck)
        tier_counts["low_confidence"] += 1

    for ck in csm_by_key:
        if ck in matched_csm:
            continue
        links_by_csm[ck] = PrefillLinkDraft(
            csm=csm_by_key[ck],
            prefill_status="NO_PRIOR_FORM_DATA",
            reason="no_purple_row",
        )
        tier_counts["no_form_data"] += 1

    links = [links_by_csm[ck] for ck in csm_by_key]
    return links, tier_counts, variant_counts, stats


# ── Dry-run orchestration + summaries ─────────────────────────────────────────


def format_csm_dry_run_summary(result: CsmDedupeResult, *, path: str | Path) -> str:
    lines: list[str] = []
    lines.append("=== Drake Prefill CSM dry-run (composite dedupe + disposition) ===")
    lines.append(f"file: {path}")
    lines.append("")
    lines.append("--- Axis: row counts ---")
    lines.append(f"rows_read (data):           {result.rows_read}")
    lines.append(f"skipped (empty name/ssn):   {result.empty_name_or_ssn_skipped}")
    lines.append(f"clients_after_dedupe:       {result.clients_after}")
    lines.append(f"groups_with_history (>1):   {result.history_groups}")
    lines.append("")
    lines.append("--- Axis: status resolution ---")
    lines.append(
        f"status_recoveries (anchor null <- earlier non-null): {len(result.recoveries)}"
    )
    for status, n in result.recovery_status_counts.most_common():
        lines.append(f"  {n:5d}  {status}")
    lines.append("")
    lines.append("disposition_status_distribution:")
    for status, n in result.disposition_counts.most_common():
        pct = 100.0 * n / max(result.clients_after, 1)
        lines.append(f"  {n:5d}  ({pct:5.1f}%)  {status}")
    return "\n".join(lines)


def run_phase2_dry_run(
    *,
    csm_path: str | Path,
    chunk_paths: list[str | Path],
) -> Phase2DryRun:
    csm = run_csm_dry_run(csm_path)
    purple, notes, row_counts, col_counts = join_purple_chunks(chunk_paths)
    name_counts = Counter(p.taxpayer_name for p in purple)
    dups = sorted(
        [(n, k) for n, k in name_counts.items() if k > 1],
        key=lambda x: (-x[1], x[0]),
    )
    links, tier_counts, variant_counts, collision_stats = match_csm_to_purple(
        csm.clients, purple
    )

    prefill_status_counts: Counter = Counter(L.prefill_status for L in links)
    no_form_by_disposition: Counter = Counter(
        L.csm["disposition_status"]
        for L in links
        if L.prefill_status == "NO_PRIOR_FORM_DATA"
    )
    purple_return_type_counts: Counter = Counter(
        (p.return_type or "(null)") for p in purple
    )
    matched_return_type_counts: Counter = Counter(
        (L.return_type or "(null)")
        for L in links
        if L.prefill_status == "PRIOR_YEAR_FORMS_AVAILABLE"
    )
    # 1040SR must survive join + match — CSM Type column reports them as 1040,
    # so the purple Return Type column is the only source of truth.
    n_sr_purple = purple_return_type_counts.get("1040SR", 0)
    n_sr_matched = matched_return_type_counts.get("1040SR", 0)
    if n_sr_purple == 0:
        raise RuntimeError(
            "Purple chunks contain no Return Type=1040SR rows — unexpected for TY2024."
        )
    notes.append(
        f"Return Type 1040SR: {n_sr_purple} purple → {n_sr_matched} PRIOR_YEAR matched"
    )

    return Phase2DryRun(
        csm=csm,
        chunk_row_counts=row_counts,
        chunk_col_counts=col_counts,
        purple_unique_names=len({p.name_norm for p in purple}),
        purple_duplicate_names=dups,
        links=links,
        tier_counts=tier_counts,
        variant_counts=variant_counts,
        prefill_status_counts=prefill_status_counts,
        no_form_by_disposition=no_form_by_disposition,
        collision_stats=collision_stats,
        purple_return_type_counts=purple_return_type_counts,
        matched_return_type_counts=matched_return_type_counts,
        assertions_ok=notes,
    )


def format_phase2_dry_run_summary(result: Phase2DryRun, *, paths: dict[str, str]) -> str:
    n_csm = result.csm.clients_after
    cs = result.collision_stats
    floor = n_csm - result.purple_unique_names
    lines: list[str] = []
    lines.append("=== Drake Prefill Phase 2 dry-run (purple + match) ===")
    for k, v in paths.items():
        lines.append(f"{k}: {v}")
    lines.append("")
    lines.append("--- Purple parse assertions ---")
    for note in result.assertions_ok:
        lines.append(f"  OK  {note}")
    lines.append(f"chunk_row_counts: {result.chunk_row_counts} (expect all 1298)")
    lines.append(f"chunk_col_counts: {result.chunk_col_counts} (expect 23,24,25)")
    lines.append(f"purple_unique_names: {result.purple_unique_names}")
    lines.append(f"purple_duplicate_names: {len(result.purple_duplicate_names)}")
    for name, k in result.purple_duplicate_names:
        lines.append(f"  {k}x  {name}")
    lines.append("")
    lines.append("--- Return Type (purple → PRIOR_YEAR matched) ---")
    lines.append(
        "(CSM Type never has 1040SR — Drake reports them as 1040; "
        "purple Return Type is authoritative)"
    )
    all_rts = sorted(
        set(result.purple_return_type_counts) | set(result.matched_return_type_counts),
        key=lambda rt: (-result.purple_return_type_counts.get(rt, 0), rt),
    )
    for rt in all_rts:
        p_n = result.purple_return_type_counts.get(rt, 0)
        m_n = result.matched_return_type_counts.get(rt, 0)
        flag = ""
        if rt == "1040SR" and m_n < p_n:
            flag = f"  ← {p_n - m_n} not yet PRIOR_YEAR (manual/low-conf)"
        lines.append(f"  purple {p_n:4d} → matched {m_n:4d}  {rt}{flag}")
    lines.append("")
    lines.append("--- Collision stats (deterministic, unique purple names) ---")
    lines.append(f"uncontested unique hit:     {cs.uncontested:5d}")
    lines.append(f"contested (>1 CSM client):  {cs.contested:5d}")
    lines.append(f"  resolved by precedence:   {cs.resolved_by_precedence:5d}")
    lines.append(f"  still tied (manual):      {cs.still_tied:5d}")
    lines.append(f"no CSM hit at all:          {cs.no_csm_hit:5d}")
    lines.append(
        f"fuzzy multi-claim groups:   {cs.fuzzy_multi_claim_groups:5d}  "
        f"(precedence={cs.fuzzy_resolved_by_precedence}, "
        f"score={cs.fuzzy_resolved_by_score}, "
        f"tied={cs.fuzzy_still_tied})"
    )
    lines.append("")
    lines.append("--- Match tiers ---")
    det = result.tier_counts.get("deterministic", 0)
    fuzz_n = result.tier_counts.get("fuzzy", 0)
    manual = result.tier_counts.get("manual", 0)
    low_conf = result.prefill_status_counts.get("LOW_CONFIDENCE_NO_MATCH", 0)
    no_form = result.prefill_status_counts.get("NO_PRIOR_FORM_DATA", 0)
    available = result.prefill_status_counts.get("PRIOR_YEAR_FORMS_AVAILABLE", 0)
    needs = result.prefill_status_counts.get("NEEDS_MANUAL_LINK", 0)
    lines.append(f"deterministic:              {det:5d}")
    lines.append(f"fuzzy (>=90):               {fuzz_n:5d}")
    lines.append(f"manual (tier):              {manual:5d}")
    lines.append(
        f"total PRIOR_YEAR_FORMS:     {available:5d}  "
        f"({100 * available / max(n_csm, 1):.1f}% of {n_csm} CSM; "
        f"{100 * available / max(result.purple_unique_names, 1):.1f}% of "
        f"{result.purple_unique_names} distinct purple)"
    )
    lines.append("")
    lines.append("--- Prefill status totals ---")
    lines.append(f"PRIOR_YEAR_FORMS_AVAILABLE: {available:5d}")
    lines.append(f"NEEDS_MANUAL_LINK:          {needs:5d}")
    lines.append(f"LOW_CONFIDENCE_NO_MATCH:    {low_conf:5d}")
    lines.append(f"NO_PRIOR_FORM_DATA:         {no_form:5d}")
    delta = no_form - floor
    lines.append(
        f"distance from {floor} floor:   {delta:+d}  "
        f"(CSM {n_csm} − distinct purple {result.purple_unique_names} = {floor}; "
        f"below floor ⇒ contested still parked in manual/low-conf)"
    )
    manual_reasons = Counter(
        L.reason for L in result.links if L.prefill_status == "NEEDS_MANUAL_LINK"
    )
    if manual_reasons:
        lines.append("NEEDS_MANUAL_LINK reason breakdown:")
        for reason, n in manual_reasons.most_common():
            lines.append(f"  {n:5d}  {reason}")
    low_reasons = Counter(
        L.reason for L in result.links if L.prefill_status == "LOW_CONFIDENCE_NO_MATCH"
    )
    if low_reasons:
        lines.append("LOW_CONFIDENCE_NO_MATCH reason breakdown:")
        for reason, n in low_reasons.most_common():
            lines.append(f"  {n:5d}  {reason}")
    lines.append("")
    lines.append("--- Variant transform hit distribution (auto-accept only) ---")
    for label, n in result.variant_counts.most_common():
        lines.append(f"  {n:5d}  {label}")
    lines.append("")
    lines.append("--- CROSS-TAB: NO_PRIOR_FORM_DATA by disposition ---")
    lines.append(
        "(should skew to PY_INCOMPLETE / PY_ROLLOVER_ONLY; "
        "large PY_FILED_ACCEPTED = match failure)"
    )
    for disp, n in result.no_form_by_disposition.most_common():
        pct = 100.0 * n / max(no_form, 1)
        flag = "  <-- investigate" if disp == "PY_FILED_ACCEPTED" and n > 40 else ""
        lines.append(f"  {n:5d}  ({pct:5.1f}%)  {disp}{flag}")
    accepted_no_form = result.no_form_by_disposition.get("PY_FILED_ACCEPTED", 0)
    if accepted_no_form > 40:
        lines.append(
            f"WARNING: {accepted_no_form} PY_FILED_ACCEPTED in NO_PRIOR_FORM_DATA — "
            "match likely failing"
        )
    else:
        lines.append(
            f"OK: PY_FILED_ACCEPTED in NO_PRIOR_FORM_DATA = {accepted_no_form} "
            "(watch for growth into the hundreds)"
        )
    return "\n".join(lines)


# ── Persist (Step 4) ──────────────────────────────────────────────────────────


def _form_counts_json(form_counts: Optional[dict[str, Any]]) -> str:
    """Serialize form counts preserving null vs key-absent (no null→0 coalesce)."""

    def _default(obj: Any) -> Any:
        if isinstance(obj, Decimal):
            return float(obj)
        raise TypeError(f"not JSON-serializable: {type(obj)!r}")

    return json.dumps(form_counts or {}, ensure_ascii=False, default=_default)


def _create_prefill_batch(
    conn: sqlite3.Connection,
    *,
    tax_year: int,
    source_paths: list[str],
    row_count: int,
) -> int:
    """Insert an import_batches row for audit; hash is unique per apply run."""
    import uuid

    stamp = utc_now()
    material = (
        f"drake_prefill|{tax_year}|{stamp}|{uuid.uuid4().hex}|{'|'.join(source_paths)}"
    )
    file_hash = hashlib.sha256(material.encode("utf-8")).hexdigest()
    filename = f"drake_prefill_ty{tax_year}.json"
    cur = conn.execute(
        """
        INSERT INTO import_batches (filename, file_hash, imported_at, status, row_count)
        VALUES (?, ?, ?, 'prefill_import', ?)
        """,
        (filename, file_hash, stamp, row_count),
    )
    return int(cur.lastrowid)


def upsert_prefill_links(
    conn: sqlite3.Connection,
    *,
    tax_year: int,
    links: list[PrefillLinkDraft],
    source_paths: list[str] | None = None,
    batch_id: Optional[int] = None,
) -> PrefillWriteStats:
    """
    Upsert drake_prefill_links + drake_form_prefill.

    Unique key: (tax_year, csm_ssn_last4, csm_name_norm).
    Never overwrite a row where match_tier = 'manual' (human adjudication).
    """
    stats = PrefillWriteStats()
    ts = utc_now()
    if batch_id is None:
        batch_id = _create_prefill_batch(
            conn,
            tax_year=tax_year,
            source_paths=source_paths or [],
            row_count=len(links),
        )
    stats.batch_id = batch_id

    for draft in links:
        c = draft.csm
        ssn = c["csm_ssn_last4"]
        name_norm = c["csm_name_norm"]
        existing = conn.execute(
            """
            SELECT id, match_tier FROM drake_prefill_links
            WHERE tax_year = ? AND csm_ssn_last4 = ? AND csm_name_norm = ?
            """,
            (tax_year, ssn, name_norm),
        ).fetchone()

        if existing is not None and existing["match_tier"] == "manual":
            stats.skipped_manual += 1
            continue

        link_vals = (
            tax_year,
            ssn,
            c["csm_name_raw"],
            name_norm,
            draft.prefill_status,
            c.get("disposition_status"),
            c.get("csm_status_raw"),
            c.get("csm_status_as_of"),
            c.get("csm_anchor_changed"),
            draft.purple_name,
            draft.match_tier,
            draft.match_score,
            draft.matched_variant,
            batch_id,
            ts,
        )

        if existing is None:
            cur = conn.execute(
                """
                INSERT INTO drake_prefill_links (
                  tax_year, csm_ssn_last4, csm_name_raw, csm_name_norm,
                  prefill_status, disposition_status,
                  csm_status_raw, csm_status_as_of, csm_anchor_changed,
                  purple_name, match_tier, match_score, matched_variant,
                  import_batch_id, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                """,
                link_vals + (ts,),
            )
            link_id = int(cur.lastrowid)
            stats.inserted += 1
        else:
            link_id = int(existing["id"])
            conn.execute(
                """
                UPDATE drake_prefill_links SET
                  csm_name_raw = ?,
                  prefill_status = ?,
                  disposition_status = ?,
                  csm_status_raw = ?,
                  csm_status_as_of = ?,
                  csm_anchor_changed = ?,
                  purple_name = ?,
                  match_tier = ?,
                  match_score = ?,
                  matched_variant = ?,
                  import_batch_id = ?,
                  updated_at = ?
                WHERE id = ?
                  AND (match_tier IS NULL OR match_tier != 'manual')
                """,
                (
                    c["csm_name_raw"],
                    draft.prefill_status,
                    c.get("disposition_status"),
                    c.get("csm_status_raw"),
                    c.get("csm_status_as_of"),
                    c.get("csm_anchor_changed"),
                    draft.purple_name,
                    draft.match_tier,
                    draft.match_score,
                    draft.matched_variant,
                    batch_id,
                    ts,
                    link_id,
                ),
            )
            stats.updated += 1

        has_forms = (
            draft.prefill_status == "PRIOR_YEAR_FORMS_AVAILABLE"
            and draft.form_counts is not None
        )
        if has_forms:
            fc_json = _form_counts_json(draft.form_counts)
            src_json = json.dumps(draft.source_files or [], ensure_ascii=False)
            conn.execute(
                """
                INSERT INTO drake_form_prefill (
                  link_id, tax_year, form_counts, return_type, source_files, created_at
                ) VALUES (?, ?, ?, ?, ?, ?)
                ON CONFLICT(link_id) DO UPDATE SET
                  tax_year = excluded.tax_year,
                  form_counts = excluded.form_counts,
                  return_type = excluded.return_type,
                  source_files = excluded.source_files
                """,
                (
                    link_id,
                    tax_year,
                    fc_json,
                    draft.return_type,
                    src_json,
                    ts,
                ),
            )
            stats.form_rows_upserted += 1
        else:
            cur = conn.execute(
                "DELETE FROM drake_form_prefill WHERE link_id = ?",
                (link_id,),
            )
            stats.form_rows_deleted += int(cur.rowcount or 0)

    conn.commit()
    return stats


def format_prefill_write_summary(stats: PrefillWriteStats) -> str:
    lines = [
        "=== Prefill DB write ===",
        f"batch_id:             {stats.batch_id}",
        f"links inserted:       {stats.inserted}",
        f"links updated:        {stats.updated}",
        f"skipped (manual):     {stats.skipped_manual}",
        f"form_counts upserted: {stats.form_rows_upserted}",
        f"form_counts deleted:  {stats.form_rows_deleted}",
    ]
    return "\n".join(lines)


def resolve_db_path(requested: str = "") -> str:
    """
    Pick a SQLite path SQLite can actually open.

    Mapped drives (``T:\\``) often fail on the app server even when Explorer
    can see them — prefer the local ``C:\\TaxOps\\taxops\\taxops.db`` or the
    UNC ``\\\\Xcel-server\\taxops\\taxops.db``.
    """
    from db import DB_PATH

    candidates: list[str] = []
    if requested:
        candidates.append(requested)
        p = Path(requested)
        # T:\taxops\foo → \\Xcel-server\taxops\foo and C:\TaxOps\taxops\foo
        parts = p.parts
        if len(parts) >= 2 and parts[0].upper().rstrip("\\") in ("T:", "T"):
            rest = Path(*parts[1:]) if len(parts) > 1 else Path()
            # parts like ('T:\\', 'taxops', 'taxops.db') or ('T:', 'taxops', ...)
            if parts[1].lower() == "taxops":
                sub = Path(*parts[2:]) if len(parts) > 2 else Path("taxops.db")
            else:
                sub = Path(*parts[1:])
            candidates.append(str(Path(r"C:\TaxOps\taxops") / sub))
            candidates.append(str(Path(r"\\Xcel-server\taxops") / sub))
    candidates.append(DB_PATH)
    candidates.append(r"C:\TaxOps\taxops\taxops.db")
    candidates.append(r"\\Xcel-server\taxops\taxops.db")

    seen: set[str] = set()
    for cand in candidates:
        if not cand or cand in seen:
            continue
        seen.add(cand)
        path = Path(cand)
        try:
            if path.is_file():
                # Probe open — catches "unable to open database file" on bad mappings
                probe = sqlite3.connect(str(path))
                probe.close()
                return str(path)
        except (OSError, sqlite3.Error):
            continue

    # Last resort: return first requested / DB_PATH for init_schema create
    return requested or DB_PATH


# ── Spouse / dependent household ingestion (Step 4 Part B) ────────────────────


@dataclass
class HouseholdRecord:
    taxpayer_name: str
    name_norm: str
    taxpayer_dob: Optional[str]
    taxpayer_phone: Optional[str]
    taxpayer_email: Optional[str]
    spouse_name: Optional[str]
    spouse_dob: Optional[str]
    spouse_phone: Optional[str]
    dependents: list[dict[str, Optional[str]]]
    source_file: str
    source_rows: int


@dataclass
class SpouseFileParseResult:
    path: str
    shape: str  # "combined" (TY2024) | "split" (TY2025)
    header_cols: int
    rows_read: int
    rows_padded: int
    rows_skipped_totals: int
    taxpayers: list[HouseholdRecord]
    inconsistency_logs: list[str]


@dataclass
class HouseholdJoinResult:
    exact_matches: int
    spouse_unmatched: list[str]
    purple_unmatched: list[str]
    attached_to_links: int
    no_link_for_match: int


@dataclass
class HouseholdWriteStats:
    upserted: int = 0
    skipped_no_link: int = 0


def _blank_to_none(value: Any) -> Optional[str]:
    if value is None:
        return None
    s = str(value).strip()
    return s if s else None


def _is_totals_name(name: str) -> bool:
    return (name or "").strip().upper().startswith("TOTALS")


def load_spouse_csv(path: str | Path) -> SpouseFileParseResult:
    """
    Parse a Drake spouse/dependent export (TY2024 combined or TY2025 split).

    Rows 1-2 title/timestamp; row 3 header. Ragged rows padded to header width;
    hard-fail if any row exceeds header. Group by taxpayer; dependents collected.
    """
    path = Path(path)
    with path.open(newline="", encoding="utf-8-sig") as f:
        rows = list(csv.reader(f))
    if len(rows) < 3:
        raise ValueError(f"{path}: expected title + timestamp + header rows")

    header = [(h or "").strip() for h in rows[2]]
    header_set = set(header)
    if "Taxpayer Name" in header_set:
        shape = "combined"
    elif "Taxpayer First Name" in header_set and "Taxpayer Last Name" in header_set:
        shape = "split"
    else:
        raise ValueError(f"{path}: unrecognized spouse header: {header}")

    width = len(header)
    data_rows = rows[3:]
    padded = 0
    skipped_totals = 0
    groups: dict[str, list[dict[str, str]]] = defaultdict(list)
    display_name: dict[str, str] = {}

    for raw in data_rows:
        if raw is None or not any((c or "").strip() for c in raw):
            continue
        if len(raw) > width:
            raise ValueError(
                f"{path}: row wider than header ({len(raw)} > {width}): {raw!r}"
            )
        if len(raw) < width:
            padded += 1
            raw = list(raw) + [""] * (width - len(raw))
        fields = {header[i]: (raw[i] if i < len(raw) else "") for i in range(width)}

        if shape == "combined":
            tp_name = (fields.get("Taxpayer Name") or "").strip()
        else:
            first = (fields.get("Taxpayer First Name") or "").strip()
            last = (fields.get("Taxpayer Last Name") or "").strip()
            tp_name = f"{first} {last}".strip() if first else last

        if not tp_name:
            continue
        if _is_totals_name(tp_name):
            skipped_totals += 1
            continue

        nrm = normalize_purple_name(tp_name)
        groups[nrm].append(fields)
        display_name[nrm] = tp_name

    taxpayers: list[HouseholdRecord] = []
    inconsistency_logs: list[str] = []

    for nrm, flist in groups.items():
        def _pick(key: str, _flist: list[dict[str, str]] = flist, _nrm: str = nrm) -> Optional[str]:
            vals = {_blank_to_none(f.get(key)) for f in _flist}
            vals.discard(None)
            if len(vals) > 1:
                inconsistency_logs.append(
                    f"{display_name.get(_nrm, _nrm)}: inconsistent {key}: {sorted(vals)}"
                )
            for f in _flist:
                v = _blank_to_none(f.get(key))
                if v is not None:
                    return v
            return None

        if shape == "combined":
            phone = _pick("Taxpayer Cell Phone")
            email = None
        else:
            phone = _pick("Taxpayer Daytime Phone")
            email = _pick("Taxpayer Email Address")

        deps: list[dict[str, Optional[str]]] = []
        for f in flist:
            df = _blank_to_none(f.get("Dependent First Name"))
            dl = _blank_to_none(f.get("Dependent Last Name"))
            if df is None and dl is None:
                continue
            deps.append({"first": df, "last": dl})

        taxpayers.append(
            HouseholdRecord(
                taxpayer_name=display_name[nrm],
                name_norm=nrm,
                taxpayer_dob=_pick("Taxpayer Date of Birth"),
                taxpayer_phone=phone,
                taxpayer_email=email,
                spouse_name=_pick("Spouse Name"),
                spouse_dob=_pick("Spouse Date of Birth"),
                spouse_phone=_pick("Spouse Daytime Phone"),
                dependents=deps,
                source_file=path.name,
                source_rows=len(flist),
            )
        )

    return SpouseFileParseResult(
        path=str(path),
        shape=shape,
        header_cols=width,
        rows_read=sum(
            1 for r in data_rows if r and any((c or "").strip() for c in r)
        ),
        rows_padded=padded,
        rows_skipped_totals=skipped_totals,
        taxpayers=taxpayers,
        inconsistency_logs=inconsistency_logs,
    )


def join_household_to_purple(
    households: list[HouseholdRecord],
    purple_name_norms: set[str],
) -> HouseholdJoinResult:
    """Join on normalized purple/spouse taxpayer name — no CSM variant pass."""
    hset = {h.name_norm for h in households}
    purple_clean = {n for n in purple_name_norms if not _is_totals_name(n)}
    exact = hset & purple_clean
    return HouseholdJoinResult(
        exact_matches=len(exact),
        spouse_unmatched=sorted(hset - purple_clean),
        purple_unmatched=sorted(purple_clean - hset),
        attached_to_links=0,
        no_link_for_match=0,
    )


def upsert_household_prefill(
    conn: sqlite3.Connection,
    *,
    tax_year: int,
    households: list[HouseholdRecord],
) -> HouseholdWriteStats:
    """
    Upsert drake_household_prefill keyed by link_id.

    Joins household.name_norm → drake_prefill_links.purple_name (normalized).
    Never writes to spouses / client_dependents.
    """
    stats = HouseholdWriteStats()
    ts = utc_now()

    by_purple: dict[str, list[int]] = defaultdict(list)
    for row in conn.execute(
        """
        SELECT id, purple_name FROM drake_prefill_links
        WHERE tax_year = ? AND purple_name IS NOT NULL AND trim(purple_name) != ''
        """,
        (tax_year,),
    ).fetchall():
        by_purple[normalize_purple_name(row["purple_name"])].append(int(row["id"]))

    for h in households:
        link_ids = by_purple.get(h.name_norm) or []
        if not link_ids:
            stats.skipped_no_link += 1
            continue
        deps_json = json.dumps(h.dependents, ensure_ascii=False)
        for link_id in link_ids:
            conn.execute(
                """
                INSERT INTO drake_household_prefill (
                  link_id, tax_year,
                  taxpayer_dob, taxpayer_phone, taxpayer_email,
                  spouse_name, spouse_dob, spouse_phone,
                  dependents_json, source_file, created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                ON CONFLICT(link_id) DO UPDATE SET
                  tax_year = excluded.tax_year,
                  taxpayer_dob = excluded.taxpayer_dob,
                  taxpayer_phone = excluded.taxpayer_phone,
                  taxpayer_email = excluded.taxpayer_email,
                  spouse_name = excluded.spouse_name,
                  spouse_dob = excluded.spouse_dob,
                  spouse_phone = excluded.spouse_phone,
                  dependents_json = excluded.dependents_json,
                  source_file = excluded.source_file,
                  updated_at = excluded.updated_at
                """,
                (
                    link_id,
                    tax_year,
                    h.taxpayer_dob,
                    h.taxpayer_phone,
                    h.taxpayer_email,
                    h.spouse_name,
                    h.spouse_dob,
                    h.spouse_phone,
                    deps_json,
                    h.source_file,
                    ts,
                    ts,
                ),
            )
            stats.upserted += 1

    conn.commit()
    return stats


def run_step4_idempotency_proof(
    *,
    links: list[PrefillLinkDraft],
    households: list[HouseholdRecord],
    tax_year: int,
    source_paths: list[str],
) -> list[str]:
    """Fresh temp DB: apply twice + manual-row immutability. Does not touch prod."""
    import shutil
    import tempfile

    from db import get_connection, init_db

    lines: list[str] = []
    tmp_dir = tempfile.mkdtemp(prefix="taxops_prefill_idem_")
    db_path = str(Path(tmp_dir) / "idem.db")
    try:
        conn = get_connection(db_path)
        init_db(conn)

        s1 = upsert_prefill_links(
            conn, tax_year=tax_year, links=links, source_paths=source_paths
        )
        h1 = upsert_household_prefill(
            conn, tax_year=tax_year, households=households
        )
        n_links_1 = conn.execute(
            "SELECT COUNT(*) AS n FROM drake_prefill_links WHERE tax_year=?",
            (tax_year,),
        ).fetchone()["n"]
        n_forms_1 = conn.execute(
            "SELECT COUNT(*) AS n FROM drake_form_prefill WHERE tax_year=?",
            (tax_year,),
        ).fetchone()["n"]
        n_hh_1 = conn.execute(
            "SELECT COUNT(*) AS n FROM drake_household_prefill WHERE tax_year=?",
            (tax_year,),
        ).fetchone()["n"]

        conn.execute(
            """
            INSERT INTO drake_prefill_links (
              tax_year, csm_ssn_last4, csm_name_raw, csm_name_norm,
              prefill_status, purple_name, match_tier, match_score, matched_variant,
              resolved_by, resolved_at, created_at, updated_at
            ) VALUES (
              ?, '0001', 'MANUAL SEED, TEST', 'MANUAL SEED, TEST',
              'NEEDS_MANUAL_LINK', 'SEED PURPLE NAME', 'manual', NULL, NULL,
              'staff_seed', '2026-01-01T00:00:00+00:00', ?, ?
            )
            """,
            (tax_year, utc_now(), utc_now()),
        )
        conn.commit()
        seed_before = dict(
            conn.execute(
                "SELECT * FROM drake_prefill_links WHERE csm_name_norm='MANUAL SEED, TEST'"
            ).fetchone()
        )

        poison = PrefillLinkDraft(
            csm={
                "csm_ssn_last4": "0001",
                "csm_name_raw": "MANUAL SEED, TEST",
                "csm_name_norm": "MANUAL SEED, TEST",
                "disposition_status": "PY_STATUS_UNKNOWN",
                "csm_status_raw": None,
                "csm_status_as_of": None,
                "csm_anchor_changed": None,
            },
            prefill_status="PRIOR_YEAR_FORMS_AVAILABLE",
            match_tier="deterministic",
            match_score=100.0,
            matched_variant="single_flip",
            purple_name="SHOULD NOT WRITE",
            form_counts={"Schedule A": 1},
            return_type="1040",
            source_files=["poison"],
        )

        s2 = upsert_prefill_links(
            conn,
            tax_year=tax_year,
            links=links + [poison],
            source_paths=source_paths,
        )
        h2 = upsert_household_prefill(
            conn, tax_year=tax_year, households=households
        )
        n_links_2 = conn.execute(
            "SELECT COUNT(*) AS n FROM drake_prefill_links WHERE tax_year=?",
            (tax_year,),
        ).fetchone()["n"]
        n_forms_2 = conn.execute(
            "SELECT COUNT(*) AS n FROM drake_form_prefill WHERE tax_year=?",
            (tax_year,),
        ).fetchone()["n"]
        n_hh_2 = conn.execute(
            "SELECT COUNT(*) AS n FROM drake_household_prefill WHERE tax_year=?",
            (tax_year,),
        ).fetchone()["n"]

        seed_after = dict(
            conn.execute(
                "SELECT * FROM drake_prefill_links WHERE csm_name_norm='MANUAL SEED, TEST'"
            ).fetchone()
        )
        conn.close()

        delta_links = n_links_2 - n_links_1
        delta_forms = n_forms_2 - n_forms_1
        delta_hh = n_hh_2 - n_hh_1

        lines.append(
            f"run1: inserted={s1.inserted} forms={s1.form_rows_upserted} "
            f"household={h1.upserted} | "
            f"counts links={n_links_1} forms={n_forms_1} household={n_hh_1}"
        )
        lines.append(
            f"run2: inserted={s2.inserted} updated={s2.updated} "
            f"skipped_manual={s2.skipped_manual} "
            f"forms_upserted={s2.form_rows_upserted} household={h2.upserted}"
        )
        lines.append(
            f"deltas: linksΔ={delta_links} (expect 1=seed only) "
            f"formsΔ={delta_forms} (expect 0) householdΔ={delta_hh} (expect 0)"
        )
        ok_idem = (
            s2.inserted == 0
            and delta_forms == 0
            and delta_hh == 0
            and delta_links == 1
            and s2.skipped_manual >= 1
        )
        lines.append(
            f"idempotency (run2 inserts=0, forms/hh unchanged): "
            f"{'PASS' if ok_idem else 'FAIL'}"
        )

        watch = (
            "purple_name",
            "resolved_by",
            "resolved_at",
            "match_tier",
            "prefill_status",
        )
        identical = all(seed_before.get(k) == seed_after.get(k) for k in watch)
        lines.append(
            f"manual immutability ({', '.join(watch)}): "
            f"{'PASS' if identical else 'FAIL'}"
        )
        if identical:
            lines.append(
                f"  seed purple_name={seed_after['purple_name']!r} "
                f"resolved_by={seed_after['resolved_by']!r} "
                f"resolved_at={seed_after['resolved_at']!r}"
            )
        else:
            for k in watch:
                lines.append(
                    f"  {k}: before={seed_before.get(k)!r} after={seed_after.get(k)!r}"
                )

        status_counts = Counter(L.prefill_status for L in links)
        total = sum(status_counts.values())
        lines.append(f"link draft total={total} (== CSM deduped clients)")
        for st, n in status_counts.most_common():
            lines.append(f"  {n:5d}  {st}")
    finally:
        shutil.rmtree(tmp_dir, ignore_errors=True)
    return lines


# Link accept for prefill→clients — local; do not change name_matcher.ACCEPT_THRESHOLD.
_PREFILL_CLIENT_LINK_ACCEPT = 90


@dataclass
class DrakeClientEmailAuditStats:
    """Repair stats for Drake household email vs client email fields."""

    scanned: int = 0
    skipped_no_drake_email: int = 0
    kept_email_match: int = 0
    kept_name_match: int = 0
    cleared_taxpayer_email: int = 0
    cleared_spouse_email: int = 0
    links_unlinked: int = 0
    skipped_manual_link: int = 0
    dry_run: bool = False
    preview: list[dict[str, Any]] = field(default_factory=list)


def _normalize_email(value: str | None) -> str:
    return (value or "").strip().lower()


def _drake_client_name_score(
    csm_name_raw: str,
    client_last: str | None,
    client_first: str | None,
) -> float:
    from name_matcher import parse_name, score_client_names_pair

    last, first = parse_name(csm_name_raw or "")
    if not last:
        return 0.0
    return float(
        score_client_names_pair(last, first, client_last or "", client_first or "")
    )


def audit_drake_client_email_links(
    conn: sqlite3.Connection,
    *,
    tax_year: int | None = None,
    dry_run: bool = False,
    preview_limit: int = 40,
) -> DrakeClientEmailAuditStats:
    """
    Ensure client email fields align with Drake household email on linked prefills.

    When Drake lists ``taxpayer_email`` and a client's taxpayer/spouse email does
    not match, clear the mismatched field(s) and unlink the prefill row from the
    client unless the Drake CSM name scores ≥ ``_PREFILL_CLIENT_LINK_ACCEPT`` against
    the client (same threshold as prefill→client linking).

    Rows with ``match_tier = 'manual'`` are never unlinked; email clears still apply
    when there is no name match.
    """
    stats = DrakeClientEmailAuditStats(dry_run=dry_run)
    has_hh = conn.execute(
        "SELECT 1 FROM sqlite_master WHERE type='table' AND name='drake_household_prefill'"
    ).fetchone()
    if not has_hh:
        return stats

    year_clause = ""
    params: list[Any] = []
    if tax_year is not None:
        year_clause = "AND l.tax_year = ?"
        params.append(int(tax_year))

    rows = conn.execute(
        f"""
        SELECT
            l.id AS link_id,
            l.tax_year,
            l.csm_name_raw,
            l.client_id,
            l.match_tier,
            c.last_name,
            c.first_name,
            c.taxpayer_email,
            c.spouse_email,
            hh.taxpayer_email AS drake_email
        FROM drake_prefill_links l
        JOIN clients c ON c.id = l.client_id
        LEFT JOIN drake_household_prefill hh ON hh.link_id = l.id
        WHERE l.client_id IS NOT NULL
          {year_clause}
        ORDER BY l.id
        """,
        tuple(params),
    ).fetchall()

    ts = utc_now()

    for row in rows:
        stats.scanned += 1
        drake_email = _normalize_email(row["drake_email"])
        if not drake_email:
            stats.skipped_no_drake_email += 1
            continue

        t_email = _normalize_email(row["taxpayer_email"])
        s_email = _normalize_email(row["spouse_email"])
        email_ok = (t_email == drake_email) or (s_email == drake_email)
        if email_ok:
            stats.kept_email_match += 1
            continue

        name_score = _drake_client_name_score(
            row["csm_name_raw"], row["last_name"], row["first_name"]
        )
        name_ok = name_score >= _PREFILL_CLIENT_LINK_ACCEPT
        if name_ok:
            stats.kept_name_match += 1
            continue

        clear_tax = bool(t_email and t_email != drake_email)
        clear_spouse = bool(s_email and s_email != drake_email)
        unlink = row["match_tier"] != "manual"

        if row["match_tier"] == "manual" and (clear_tax or clear_spouse):
            stats.skipped_manual_link += 1

        if dry_run:
            if len(stats.preview) < preview_limit:
                stats.preview.append(
                    {
                        "link_id": int(row["link_id"]),
                        "client_id": int(row["client_id"]),
                        "drake_email": drake_email,
                        "taxpayer_email": t_email or None,
                        "spouse_email": s_email or None,
                        "name_score": round(name_score, 1),
                        "clear_taxpayer": clear_tax,
                        "clear_spouse": clear_spouse,
                        "unlink": unlink,
                    }
                )
            if clear_tax:
                stats.cleared_taxpayer_email += 1
            if clear_spouse:
                stats.cleared_spouse_email += 1
            if unlink:
                stats.links_unlinked += 1
            continue

        if clear_tax:
            conn.execute(
                """
                UPDATE clients
                SET taxpayer_email = NULL, updated_at = ?
                WHERE id = ? AND lower(trim(taxpayer_email)) != ?
                """,
                (ts, int(row["client_id"]), drake_email),
            )
            stats.cleared_taxpayer_email += 1
        if clear_spouse:
            conn.execute(
                """
                UPDATE clients
                SET spouse_email = NULL, updated_at = ?
                WHERE id = ? AND lower(trim(spouse_email)) != ?
                """,
                (ts, int(row["client_id"]), drake_email),
            )
            stats.cleared_spouse_email += 1
        if unlink:
            conn.execute(
                """
                UPDATE drake_prefill_links
                SET client_id = NULL, updated_at = ?
                WHERE id = ? AND (match_tier IS NULL OR match_tier != 'manual')
                """,
                (ts, int(row["link_id"])),
            )
            stats.links_unlinked += 1

    return stats


@dataclass
class HouseholdEmailBackfillStats:
    csv_with_email: int = 0
    matched_links: int = 0
    updated_rows: int = 0
    skipped_no_link: int = 0
    dry_run: bool = False


def backfill_household_emails_from_spouse_csv(
    conn: sqlite3.Connection,
    csv_path: str | Path,
    *,
    link_tax_year: int = 2024,
    dry_run: bool = False,
) -> HouseholdEmailBackfillStats:
    """
    Merge ``Taxpayer Email Address`` from a split spouse export (e.g. TY2025 CSV)
    into existing ``drake_household_prefill`` rows keyed by purple name.

    TY2024 combined spouse CSV has no email column — use this before email audit.
    """
    stats = HouseholdEmailBackfillStats(dry_run=dry_run)
    parsed = load_spouse_csv(csv_path)

    by_purple: dict[str, list[int]] = defaultdict(list)
    for row in conn.execute(
        """
        SELECT id, purple_name FROM drake_prefill_links
        WHERE tax_year = ? AND purple_name IS NOT NULL AND trim(purple_name) != ''
        """,
        (link_tax_year,),
    ).fetchall():
        by_purple[normalize_purple_name(row["purple_name"])].append(int(row["id"]))

    ts = utc_now()
    for h in parsed.taxpayers:
        email = (h.taxpayer_email or "").strip()
        if not email:
            continue
        stats.csv_with_email += 1
        link_ids = by_purple.get(h.name_norm) or []
        if not link_ids:
            stats.skipped_no_link += 1
            continue
        for link_id in link_ids:
            stats.matched_links += 1
            if dry_run:
                stats.updated_rows += 1
                continue
            cur = conn.execute(
                """
                UPDATE drake_household_prefill
                SET taxpayer_email = ?, updated_at = ?
                WHERE link_id = ?
                  AND (taxpayer_email IS NULL OR trim(taxpayer_email) = '')
                """,
                (email, ts, link_id),
            )
            if cur.rowcount:
                stats.updated_rows += 1

    return stats


def format_household_email_backfill_summary(stats: HouseholdEmailBackfillStats) -> str:
    mode = "DRY-RUN" if stats.dry_run else "APPLY"
    return "\n".join(
        [
            f"=== Household email backfill ({mode}) ===",
            f"CSV rows with email:   {stats.csv_with_email}",
            f"matched prefill links: {stats.matched_links}",
            f"household rows updated:{stats.updated_rows}",
            f"CSV names w/o link:    {stats.skipped_no_link}",
        ]
    )


def format_drake_client_email_audit_summary(stats: DrakeClientEmailAuditStats) -> str:
    mode = "DRY-RUN" if stats.dry_run else "APPLY"
    lines = [
        f"=== Drake <-> client email audit ({mode}) ===",
        f"linked rows scanned:     {stats.scanned}",
        f"skipped (no Drake email): {stats.skipped_no_drake_email}",
        f"kept (email matches):     {stats.kept_email_match}",
        f"kept (name match):        {stats.kept_name_match}",
        f"taxpayer emails cleared:  {stats.cleared_taxpayer_email}",
        f"spouse emails cleared:    {stats.cleared_spouse_email}",
        f"prefill links unlinked:   {stats.links_unlinked}",
        f"manual links (no unlink): {stats.skipped_manual_link}",
    ]
    if stats.preview:
        lines.append("preview:")
        for p in stats.preview[:15]:
            lines.append(
                f"  link {p['link_id']} client {p['client_id']} "
                f"score={p['name_score']} unlink={p['unlink']} "
                f"clear_tax={p['clear_taxpayer']} clear_spouse={p['clear_spouse']}"
            )
    if stats.scanned == 0:
        lines.append(
            "hint: no linked prefills for this tax year — prefill links are TY2024 "
            "(use --tax-year 2024, not 2025)."
        )
    elif stats.scanned > 0 and stats.skipped_no_drake_email == stats.scanned:
        lines.append(
            "hint: linked rows exist but Drake household email is empty — run with "
            "--backfill-emails-from CSVFILES/TAXPAYERspouse25.csv first "
            "(TY2024 spouse CSV has no email column)."
        )
    return "\n".join(lines)


@dataclass
class PrefillClientLinkStats:
    linked_existing: int = 0
    created_clients: int = 0
    spouse_filled: int = 0
    skipped_ambiguous: int = 0
    skipped_email_mismatch: int = 0
    already_linked: int = 0
    would_create: int = 0
    would_link_existing: int = 0
    dry_run: bool = False
    create_preview: list = field(default_factory=list)


def _prefill_link_email_compatible(
    conn: sqlite3.Connection,
    link_id: int,
    client_id: int,
    csm_name_raw: str,
) -> bool:
    """True when client emails match Drake household email or names score high enough."""
    hh = conn.execute(
        "SELECT taxpayer_email FROM drake_household_prefill WHERE link_id = ?",
        (link_id,),
    ).fetchone()
    drake_email = _normalize_email(hh["taxpayer_email"] if hh else None)
    if not drake_email:
        return True
    client = conn.execute(
        "SELECT last_name, first_name, taxpayer_email, spouse_email FROM clients WHERE id = ?",
        (client_id,),
    ).fetchone()
    if not client:
        return False
    t_email = _normalize_email(client["taxpayer_email"])
    s_email = _normalize_email(client["spouse_email"])
    if t_email == drake_email or s_email == drake_email:
        return True
    score = _drake_client_name_score(
        csm_name_raw, client["last_name"], client["first_name"]
    )
    return score >= _PREFILL_CLIENT_LINK_ACCEPT


def _spouse_from_csm_name(csm_name_raw: str) -> tuple[Optional[str], Optional[str]]:
    """Return (spouse_first, spouse_last) from a CSM ``LAST, FIRST & SPOUSE`` line."""
    from name_matcher import spouse_parts_from_display_line

    parts = spouse_parts_from_display_line(csm_name_raw)
    if not parts:
        return None, None
    sp_first, sp_rest = parts
    if not sp_first:
        return None, None
    # Shared-surname joint: spouse chunk is just a first name → use primary last
    if not sp_rest:
        from name_matcher import parse_name

        last, _ = parse_name(csm_name_raw)
        return sp_first, last or None
    # ``MARTINEZ, MARIA`` style chunk (rare in spouse_parts) — rest is leftover tokens
    return sp_first, sp_rest


def link_prefill_clients(
    conn: sqlite3.Connection,
    *,
    tax_year: int,
    statuses: tuple[str, ...] = ("PRIOR_YEAR_FORMS_AVAILABLE",),
    dry_run: bool = False,
) -> PrefillClientLinkStats:
    """
    Attach ``drake_prefill_links.client_id`` so intake search can find prior-year
    Drake clients (incl. joint filers / 1040SR).

    - Match existing clients by ssn_last4+name or fuzzy name (≥90).
    - Create a client stub when none match (PRIOR_YEAR only by default).
    - Fill empty spouse_* from the CSM joint name (never overwrite).

    Wave 3: ``dry_run=True`` previews creates/links and writes nothing.
    """
    from name_matcher import find_client, parse_name, score_client_names_pair

    stats = PrefillClientLinkStats(dry_run=dry_run)
    ts = utc_now()
    placeholders = ",".join("?" * len(statuses))
    rows = conn.execute(
        f"""
        SELECT id, csm_ssn_last4, csm_name_raw, csm_name_norm, client_id, prefill_status
        FROM drake_prefill_links
        WHERE tax_year = ?
          AND prefill_status IN ({placeholders})
        ORDER BY id
        """,
        (tax_year, *statuses),
    ).fetchall()

    # Rebuild cache as we create clients
    client_cache = None

    for row in rows:
        if row["client_id"] is not None:
            stats.already_linked += 1
            if not dry_run:
                _fill_spouse_if_empty(
                    conn, int(row["client_id"]), row["csm_name_raw"], stats, ts
                )
            continue

        last, first = parse_name(row["csm_name_raw"] or "")
        if not last:
            stats.skipped_ambiguous += 1
            continue

        ssn = (row["csm_ssn_last4"] or "").strip() or None
        client_id: Optional[int] = None

        # 1) ssn_last4 + strong name score
        if ssn:
            cands = conn.execute(
                "SELECT id, last_name, first_name FROM clients WHERE ssn_last4 = ?",
                (ssn,),
            ).fetchall()
            scored = []
            for c in cands:
                sc = score_client_names_pair(
                    last, first, c["last_name"], c["first_name"]
                )
                if sc >= _PREFILL_CLIENT_LINK_ACCEPT:
                    scored.append((sc, int(c["id"])))
            scored.sort(reverse=True)
            if len(scored) == 1 or (len(scored) > 1 and scored[0][0] > scored[1][0]):
                client_id = scored[0][1]
            elif len(scored) > 1:
                stats.skipped_ambiguous += 1
                continue

        # 2) fuzzy / exact name match
        if client_id is None:
            match = find_client(conn, last, first, cache=client_cache)
            if (
                match
                and match["score"] >= _PREFILL_CLIENT_LINK_ACCEPT
                and not match.get("needs_review")
            ):
                client_id = int(match["client_id"])
                # If client has a different ssn, don't force-link
                if ssn:
                    existing_ssn = conn.execute(
                        "SELECT ssn_last4 FROM clients WHERE id = ?",
                        (client_id,),
                    ).fetchone()
                    other = (existing_ssn["ssn_last4"] or "").strip() if existing_ssn else ""
                    if other and other != ssn:
                        client_id = None

        # 3) create stub
        if client_id is None:
            sp_first, sp_last = _spouse_from_csm_name(row["csm_name_raw"] or "")
            if dry_run:
                stats.would_create += 1
                if len(stats.create_preview) < 25:
                    stats.create_preview.append(
                        {
                            "link_id": int(row["id"]),
                            "last": last,
                            "first": first,
                            "ssn_last4": ssn,
                        }
                    )
                continue
            cur = conn.execute(
                """
                INSERT INTO clients (
                  last_name, first_name, ssn_last4,
                  spouse_first_name, spouse_last_name,
                  created_at, updated_at
                ) VALUES (?, ?, ?, ?, ?, ?, ?)
                """,
                (last, first, ssn, sp_first, sp_last, ts, ts),
            )
            client_id = int(cur.lastrowid)
            stats.created_clients += 1
            client_cache = None  # invalidate
            if sp_first:
                stats.spouse_filled += 1
        else:
            if dry_run:
                stats.would_link_existing += 1
                continue
            stats.linked_existing += 1
            _fill_spouse_if_empty(conn, client_id, row["csm_name_raw"], stats, ts)
            # Backfill ssn_last4 when client is missing it
            if ssn:
                conn.execute(
                    """
                    UPDATE clients SET ssn_last4 = ?, updated_at = ?
                    WHERE id = ? AND (ssn_last4 IS NULL OR trim(ssn_last4) = '')
                    """,
                    (ssn, ts, client_id),
                )

        if not _prefill_link_email_compatible(
            conn, int(row["id"]), client_id, row["csm_name_raw"] or ""
        ):
            stats.skipped_email_mismatch += 1
            continue

        conn.execute(
            "UPDATE drake_prefill_links SET client_id = ?, updated_at = ? WHERE id = ?",
            (client_id, ts, row["id"]),
        )

    if not dry_run:
        audit_drake_client_email_links(conn, tax_year=tax_year, dry_run=False)
        conn.commit()
    return stats


def _fill_spouse_if_empty(
    conn: sqlite3.Connection,
    client_id: int,
    csm_name_raw: str,
    stats: PrefillClientLinkStats,
    ts: str,
) -> None:
    """Wave 4: prefer ``spouses`` upsert; also fill empty clients.spouse_* for read-through."""
    sp_first, sp_last = _spouse_from_csm_name(csm_name_raw)
    if not sp_first:
        return
    sp_row = conn.execute(
        "SELECT first_name FROM spouses WHERE client_id = ?", (client_id,)
    ).fetchone()
    if not sp_row:
        conn.execute(
            """
            INSERT INTO spouses (client_id, first_name, last_name, source, created_at)
            VALUES (?, ?, ?, 'prefill_link', ?)
            """,
            (client_id, sp_first, sp_last, ts),
        )
        stats.spouse_filled += 1
    row = conn.execute(
        "SELECT spouse_first_name, spouse_last_name FROM clients WHERE id = ?",
        (client_id,),
    ).fetchone()
    if not row:
        return
    has_sp = bool((row["spouse_first_name"] or "").strip())
    if has_sp:
        return
    conn.execute(
        """
        UPDATE clients
        SET spouse_first_name = ?, spouse_last_name = ?, updated_at = ?
        WHERE id = ?
        """,
        (sp_first, sp_last, ts, client_id),
    )
    if sp_row:
        stats.spouse_filled += 1


def format_prefill_client_link_summary(stats: PrefillClientLinkStats) -> str:
    lines = [
        "=== Prefill → clients link ===",
        f"linked to existing:   {stats.linked_existing}",
        f"clients created:      {stats.created_clients}",
        f"spouse fields filled: {stats.spouse_filled}",
        f"already linked:       {stats.already_linked}",
        f"skipped email mismatch: {stats.skipped_email_mismatch}",
        f"skipped ambiguous:    {stats.skipped_ambiguous}",
    ]
    if stats.dry_run:
        lines.extend(
            [
                f"DRY-RUN would create: {stats.would_create}",
                f"DRY-RUN would link:   {stats.would_link_existing}",
                "No writes committed. Re-run with --link-clients --apply to mint.",
            ]
        )
        for p in stats.create_preview[:10]:
            lines.append(
                f"  preview create: {p.get('last')}, {p.get('first')} "
                f"ssn={p.get('ssn_last4')} link_id={p.get('link_id')}"
            )
    return "\n".join(lines)


def main(argv: list[str] | None = None) -> int:
    parser = argparse.ArgumentParser(description="Drake prefill importer")
    parser.add_argument(
        "--csm",
        default=r"C:\TaxOps\taxops\CSVFILES\2024 CLIENTS.xlsx",
    )
    parser.add_argument(
        "--chunk1",
        default=r"C:\TaxOps\taxops\CSVFILES\TY2024S.csv",
    )
    parser.add_argument(
        "--chunk2",
        default=r"C:\TaxOps\taxops\CSVFILES\TY2024S2439-5405.csv",
    )
    parser.add_argument(
        "--chunk3",
        default=r"C:\TaxOps\taxops\CSVFILES\TY2024S8867-end.csv",
    )
    parser.add_argument(
        "--spouse2024",
        default=r"C:\TaxOps\taxops\CSVFILES\TY2024Spouses.csv",
    )
    parser.add_argument(
        "--spouse2025",
        default=r"C:\TaxOps\taxops\CSVFILES\TAXPAYERspouse25.csv",
    )
    parser.add_argument("--tax-year", type=int, default=2024)
    parser.add_argument(
        "--phase",
        choices=("1", "2", "4"),
        default="4",
        help="1=CSM only, 2=purple+match, 4=Step4 dry-run (links+household+idempotency)",
    )
    parser.add_argument(
        "--dry-run",
        action="store_true",
        default=True,
        help="Parse + match only; do not write (default)",
    )
    parser.add_argument(
        "--apply",
        action="store_true",
        help="Upsert drake_prefill_links + drake_form_prefill (implies not dry-run)",
    )
    parser.add_argument(
        "--db",
        default="",
        help="SQLite path for --apply (default: TAXOPS_DB / config.DB_PATH)",
    )
    parser.add_argument(
        "--init-schema",
        action="store_true",
        help="Run init_db/_migrate on --db before apply (needed for fresh copies)",
    )
    parser.add_argument(
        "--link-clients",
        action="store_true",
        help=(
            "Link/create TaxOps clients for PRIOR_YEAR prefill rows. "
            "Alone = dry-run preview (Wave 3). With --apply = mint stubs."
        ),
    )
    parser.add_argument(
        "--audit-emails",
        action="store_true",
        help=(
            "Audit linked Drake prefills vs client email fields; clear mismatches "
            "and unlink when Drake email does not match and names do not score ≥90. "
            "Use --tax-year 2024 (prefill links are TY2024). "
            "Alone = dry-run; with --apply = write fixes."
        ),
    )
    parser.add_argument(
        "--backfill-emails-from",
        default="",
        metavar="CSV",
        help=(
            "Before --audit-emails, merge Taxpayer Email Address from a split spouse "
            "CSV (e.g. CSVFILES/TAXPAYERspouse25.csv) into drake_household_prefill."
        ),
    )
    parser.add_argument("--write-json", default="")
    parser.add_argument("-v", "--verbose", action="store_true")
    args = parser.parse_args(argv)

    logging.basicConfig(
        level=logging.DEBUG if args.verbose else logging.INFO,
        format="%(levelname)s %(message)s",
    )

    if args.phase == "1":
        result = run_csm_dry_run(args.csm)
        print(format_csm_dry_run_summary(result, path=args.csm))
        return 0

    # Standalone client-link (no CSM/purple re-parse).
    # Wave 3: --link-clients alone is dry-run; mint only with --apply.
    if args.audit_emails:
        from db import get_connection

        db_path = resolve_db_path(args.db)
        mode = "APPLY" if args.apply else "DRY-RUN"
        print(f"Auditing Drake <-> client emails ({mode}) -> {db_path}")
        conn = get_connection(db_path)
        try:
            if args.init_schema:
                from db import init_db

                init_db(db_path)
            if args.backfill_emails_from:
                # Write emails to conn so audit can see them; rollback unless --apply.
                bf_stats = backfill_household_emails_from_spouse_csv(
                    conn,
                    args.backfill_emails_from,
                    link_tax_year=args.tax_year,
                    dry_run=False,
                )
                print(format_household_email_backfill_summary(
                    HouseholdEmailBackfillStats(
                        csv_with_email=bf_stats.csv_with_email,
                        matched_links=bf_stats.matched_links,
                        updated_rows=bf_stats.updated_rows,
                        skipped_no_link=bf_stats.skipped_no_link,
                        dry_run=not args.apply,
                    )
                ))
            audit_stats = audit_drake_client_email_links(
                conn,
                tax_year=args.tax_year,
                dry_run=not args.apply,
            )
            if args.apply:
                conn.commit()
            else:
                conn.rollback()
            print(format_drake_client_email_audit_summary(audit_stats))
        finally:
            conn.close()
        return 0

    if args.link_clients:
        from db import get_connection

        db_path = resolve_db_path(args.db)
        mode = "APPLY" if args.apply else "DRY-RUN"
        print(f"Linking prefill → clients ({mode}) → {db_path}")
        conn = get_connection(db_path)
        try:
            link_stats = link_prefill_clients(
                conn, tax_year=args.tax_year, dry_run=not args.apply
            )
            print(format_prefill_client_link_summary(link_stats))
        finally:
            conn.close()
        return 0

    p2 = run_phase2_dry_run(
        csm_path=args.csm,
        chunk_paths=[args.chunk1, args.chunk2, args.chunk3],
    )
    paths = {
        "csm": args.csm,
        "chunk1": args.chunk1,
        "chunk2": args.chunk2,
        "chunk3": args.chunk3,
    }
    print(format_phase2_dry_run_summary(p2, paths=paths))
    print()
    print(format_csm_dry_run_summary(p2.csm, path=args.csm))

    # Assert link draft total == CSM deduped clients
    if len(p2.links) != p2.csm.clients_after:
        raise RuntimeError(
            f"Link draft count {len(p2.links)} != CSM clients {p2.csm.clients_after}"
        )

    spouse24 = spouse25 = None
    join24 = None
    if args.phase == "4" or args.apply:
        print()
        print("=== Step 4 — spouse/dependent parse ===")
        spouse24 = load_spouse_csv(args.spouse2024)
        spouse25 = load_spouse_csv(args.spouse2025)
        for label, sp in (("TY2024", spouse24), ("TY2025", spouse25)):
            n_deps = sum(len(t.dependents) for t in sp.taxpayers)
            multi = sum(1 for t in sp.taxpayers if t.source_rows > 1)
            print(
                f"{label}: shape={sp.shape} header_cols={sp.header_cols} "
                f"rows_read={sp.rows_read} padded={sp.rows_padded} "
                f"skipped_totals={sp.rows_skipped_totals}"
            )
            print(
                f"  taxpayers_grouped={len(sp.taxpayers)}  "
                f"dependents_collected={n_deps}  "
                f"taxpayers_with_>1_row={multi}"
            )
            if sp.inconsistency_logs:
                print(f"  inconsistencies: {len(sp.inconsistency_logs)}")
                for msg in sp.inconsistency_logs[:10]:
                    print(f"    {msg}")

        # Use purple file norms for the 1295-vs-1295 audit
        purple_joined, _, _, _ = join_purple_chunks(
            [args.chunk1, args.chunk2, args.chunk3]
        )
        purple_file_norms = {p.name_norm for p in purple_joined}
        join24 = join_household_to_purple(spouse24.taxpayers, purple_file_norms)
        print()
        print("--- Spouse↔purple join (TY2024) ---")
        print(
            f"exact matches: {join24.exact_matches}  "
            f"(expect 1294 with TOTALS, or 1294–1295 after TOTALS skip)"
        )
        print(f"spouse unmatched: {len(join24.spouse_unmatched)}")
        for n in join24.spouse_unmatched[:5]:
            print(f"  spouse-only: {n}")
        print(f"purple unmatched: {len(join24.purple_unmatched)}")
        for n in join24.purple_unmatched[:5]:
            print(f"  purple-only: {n}")

        # How many exact matches have a PRIOR_YEAR / any link with that purple_name
        link_purple = {
            normalize_purple_name(L.purple_name)
            for L in p2.links
            if L.purple_name
        }
        attachable = sum(1 for h in spouse24.taxpayers if h.name_norm in link_purple)
        print(f"household rows attachable to a link purple_name: {attachable}")

        print()
        print("=== Step 4 — idempotency proof (temp DB, not prod) ===")
        proof = run_step4_idempotency_proof(
            links=p2.links,
            households=spouse24.taxpayers,
            tax_year=args.tax_year,
            source_paths=list(paths.values()) + [args.spouse2024],
        )
        for line in proof:
            print(line)

        if not args.apply:
            print()
            print(
                "DRY-RUN ONLY — prod not written. "
                "Re-run with --apply --init-schema after greenlight "
                "(omit --db on the server so it uses C:\\TaxOps\\taxops\\taxops.db; "
                "do not pass T:\\ — SQLite often cannot open mapped drives)."
            )

    if args.apply:
        from db import get_connection, get_schema_version, init_db

        db_path = resolve_db_path(args.db)
        print()
        print(f"Applying prefill upsert → {db_path}")
        try:
            conn = get_connection(db_path)
        except sqlite3.OperationalError as exc:
            print(
                f"ERROR: cannot open database {db_path!r}: {exc}\n"
                "On the app server use the local path, e.g.\n"
                '  --db "C:\\TaxOps\\taxops\\taxops.db"\n'
                "or omit --db. Avoid T:\\ mapped-drive paths."
            )
            return 2
        try:
            if args.init_schema:
                init_db(conn)
            ver = get_schema_version(conn)
            tables = {
                r[0]
                for r in conn.execute(
                    "SELECT name FROM sqlite_master WHERE type='table'"
                ).fetchall()
            }
            if "drake_prefill_links" not in tables:
                print(
                    "ERROR: drake_prefill_links missing. "
                    "Re-run with --init-schema on a DB copy, or migrate to schema v24 first."
                )
                return 2
            if "drake_household_prefill" not in tables:
                print("ERROR: drake_household_prefill missing — run with --init-schema.")
                return 2
            if ver < 24:
                print(
                    f"WARNING: schema_version={ver} (<24). "
                    "Tables exist but version stamp is stale; continuing."
                )
            if len(p2.links) != p2.csm.clients_after:
                raise RuntimeError(
                    f"Refusing apply: links {len(p2.links)} != CSM {p2.csm.clients_after}"
                )
            write_stats = upsert_prefill_links(
                conn,
                tax_year=args.tax_year,
                links=p2.links,
                source_paths=list(paths.values()),
            )
            print(format_prefill_write_summary(write_stats))
            if spouse24 is None:
                spouse24 = load_spouse_csv(args.spouse2024)
            hh_stats = upsert_household_prefill(
                conn, tax_year=args.tax_year, households=spouse24.taxpayers
            )
            print(
                f"household upserted={hh_stats.upserted} "
                f"skipped_no_link={hh_stats.skipped_no_link}"
            )
            # Post-write verification
            n_links = conn.execute(
                "SELECT COUNT(*) AS n FROM drake_prefill_links WHERE tax_year = ?",
                (args.tax_year,),
            ).fetchone()["n"]
            n_forms = conn.execute(
                "SELECT COUNT(*) AS n FROM drake_form_prefill WHERE tax_year = ?",
                (args.tax_year,),
            ).fetchone()["n"]
            n_hh = conn.execute(
                "SELECT COUNT(*) AS n FROM drake_household_prefill WHERE tax_year = ?",
                (args.tax_year,),
            ).fetchone()["n"]
            by_status = conn.execute(
                """
                SELECT prefill_status, COUNT(*) AS n
                FROM drake_prefill_links WHERE tax_year = ?
                GROUP BY prefill_status ORDER BY n DESC
                """,
                (args.tax_year,),
            ).fetchall()
            by_rt = conn.execute(
                """
                SELECT return_type, COUNT(*) AS n
                FROM drake_form_prefill WHERE tax_year = ?
                GROUP BY return_type ORDER BY n DESC
                """,
                (args.tax_year,),
            ).fetchall()
            print("--- Post-write verification ---")
            print(f"links in DB (ty{args.tax_year}): {n_links}")
            print(f"form_prefill rows:             {n_forms}")
            print(f"household_prefill rows:        {n_hh}")
            if n_links != p2.csm.clients_after:
                raise RuntimeError(
                    f"Post-write link count {n_links} != CSM {p2.csm.clients_after}"
                )
            for row in by_status:
                print(f"  {row['n']:5d}  {row['prefill_status']}")
            print("form_prefill return_type:")
            for row in by_rt:
                mark = "  ← includes senior 1040" if row["return_type"] == "1040SR" else ""
                print(f"  {row['n']:5d}  {row['return_type']!r}{mark}")
            # Always wire clients after apply so intake search can find them
            link_stats = link_prefill_clients(conn, tax_year=args.tax_year)
            print(format_prefill_client_link_summary(link_stats))
        finally:
            conn.close()

    if args.write_json:
        out = Path(args.write_json)
        out.parent.mkdir(parents=True, exist_ok=True)
        payload = {
            "tax_year": args.tax_year,
            "tier_counts": dict(p2.tier_counts),
            "variant_counts": dict(p2.variant_counts),
            "prefill_status_counts": dict(p2.prefill_status_counts),
            "no_form_by_disposition": dict(p2.no_form_by_disposition),
            "collision_stats": {
                "uncontested": p2.collision_stats.uncontested,
                "contested": p2.collision_stats.contested,
                "resolved_by_precedence": p2.collision_stats.resolved_by_precedence,
                "still_tied": p2.collision_stats.still_tied,
                "no_csm_hit": p2.collision_stats.no_csm_hit,
                "fuzzy_multi_claim_groups": p2.collision_stats.fuzzy_multi_claim_groups,
                "fuzzy_resolved_by_precedence": (
                    p2.collision_stats.fuzzy_resolved_by_precedence
                ),
                "fuzzy_resolved_by_score": p2.collision_stats.fuzzy_resolved_by_score,
                "fuzzy_still_tied": p2.collision_stats.fuzzy_still_tied,
            },
            "purple_duplicate_names": p2.purple_duplicate_names,
            "links": [
                {
                    "csm_ssn_last4": L.csm["csm_ssn_last4"],
                    "csm_name_raw": L.csm["csm_name_raw"],
                    "disposition_status": L.csm["disposition_status"],
                    "prefill_status": L.prefill_status,
                    "match_tier": L.match_tier,
                    "match_score": L.match_score,
                    "matched_variant": L.matched_variant,
                    "purple_name": L.purple_name,
                    "return_type": L.return_type,
                    "reason": L.reason,
                    "form_counts": L.form_counts,
                }
                for L in p2.links
            ],
        }
        out.write_text(json.dumps(payload, indent=2), encoding="utf-8")
        print(f"\nwrote {out}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
