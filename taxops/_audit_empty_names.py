"""Deep dive: 394 empty-name CSM skips + SSN collisions + prod prefill gap."""
from collections import Counter, defaultdict
from openpyxl import load_workbook
from drake_prefill_importer import (
    load_csm_xlsx, composite_dedupe_csm, extract_ssn_last4, _cell_str, normalize_csm_name
)
from db import get_connection, get_schema_version

CSM = r"T:\taxops\CSVFILES\2024 CLIENTS.xlsx"

rows = load_csm_xlsx(CSM)
empty_name = [r for r in rows if not r.name_norm]
print(f"empty-name rows: {len(empty_name)}")
print(f"unique ssn among empty-name: {len({r.ssn_last4 for r in empty_name})}")
print(f"status among empty-name: {Counter(r.status or '(null)' for r in empty_name).most_common()}")
print(f"return_type among empty-name: {Counter(r.return_type or '(null)' for r in empty_name).most_common(10)}")

# Do empty-name SSNs also appear with a named row?
named = [r for r in rows if r.name_norm]
named_ssns = {r.ssn_last4 for r in named}
empty_ssns = {r.ssn_last4 for r in empty_name}
only_empty = empty_ssns - named_ssns
also_named = empty_ssns & named_ssns
print(f"\nempty-name SSNs that ALSO have a named row: {len(also_named)}")
print(f"empty-name SSNs with NO named row anywhere: {len(only_empty)}")

# After dedupe, how many of also_named are in the 1401?
dedupe = composite_dedupe_csm(rows)
kept_ssns = {c["csm_ssn_last4"] for c in dedupe.clients}
print(f"of also_named, kept in 1401 via named row: {len(also_named & kept_ssns)}")
print(f"only_empty SSNs that somehow kept: {len(only_empty & kept_ssns)}")

print("\n--- Sample only_empty (SSN never has a Client Name) up to 40 ---")
by_ssn = defaultdict(list)
for r in empty_name:
    if r.ssn_last4 in only_empty:
        by_ssn[r.ssn_last4].append(r)
for i, (ssn, grp) in enumerate(sorted(by_ssn.items())[:40]):
    g = sorted(grp, key=lambda x: x.last_change, reverse=True)[0]
    print(
        f"  {ssn}  type={g.return_type!r} status={g.status!r} "
        f"started={g.started!r} completed={g.completed!r} "
        f"changed={g.last_change_raw!r} by={g.changed_by!r}"
    )

# Raw xlsx: confirm what "empty name" looks like — maybe spaces / special?
wb = load_workbook(CSM, read_only=True, data_only=True)
ws = wb[wb.sheetnames[0]]
rows_iter = ws.iter_rows(values_only=True)
next(rows_iter); next(rows_iter)
headers = [str(h).strip() if h is not None else "" for h in next(rows_iter)]
hu = {h.upper(): h for h in headers}
name_h = hu["CLIENT NAME"]
id_h = hu["ID (LAST 4)"]
weird = []
for vals in rows_iter:
    if vals is None:
        continue
    row = {headers[i]: vals[i] if i < len(vals) else None for i in range(len(headers))}
    raw_name = row.get(name_h)
    ssn = extract_ssn_last4(row.get(id_h))
    if ssn and (raw_name is None or str(raw_name).strip() == ""):
        weird.append((ssn, raw_name, row.get(hu.get("STATUS")), row.get(hu.get("TYPE"))))
wb.close()
print(f"\nraw blank-name rows recounted: {len(weird)}")
print(f"types: {Counter(t for _,_,_,t in weird).most_common()}")

# Prod DB: confirm prefill empty + when schema became 24
conn = get_connection(r"T:\taxops\taxops.db")
print(f"\n=== live taxops.db ===")
print(f"schema {get_schema_version(conn)}")
print("prefill", conn.execute(
    "SELECT COUNT(*) n FROM sqlite_master WHERE name='drake_prefill_links'"
).fetchone()["n"])
try:
    print("link count", conn.execute("SELECT COUNT(*) n FROM drake_prefill_links").fetchone()["n"])
except Exception as e:
    print("links err", e)

# Total clients / returns by year
years = conn.execute(
    "SELECT tax_year, COUNT(*) n, COUNT(DISTINCT client_id) c FROM returns GROUP BY tax_year ORDER BY tax_year"
).fetchall()
print("returns by year:")
for y in years:
    print(f"  {y['tax_year']}: returns={y['n']} clients={y['c']}")
print("total clients", conn.execute("SELECT COUNT(*) n FROM clients").fetchone()["n"])
conn.close()

# Arithmetic
print("\n=== Coverage arithmetic ===")
print(f"CSM data rows:           {len(rows)}")
print(f"blank Client Name skips: {len(empty_name)}  (all have SSN)")
print(f"  → recoverable via named twin: {len(also_named)} SSNs ({sum(1 for r in empty_name if r.ssn_last4 in also_named)} rows)")
print(f"  → truly nameless (lost):      {len(only_empty)} SSNs ({sum(1 for r in empty_name if r.ssn_last4 in only_empty)} rows)")
print(f"meaningful rows before dedupe: {len(rows)-len(empty_name)}")
print(f"after composite dedupe:        {dedupe.clients_after}")
print(f"collapsed by history:          {len(rows)-len(empty_name) - dedupe.clients_after}")
