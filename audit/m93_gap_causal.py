"""M9.3 — multi-label causal attributes for unmatched TY2025 returns."""

from __future__ import annotations

import json
import sqlite3
from collections import Counter, defaultdict
from pathlib import Path
from typing import Any, Optional

import openpyxl

from audit import config
from audit.db import connect_audit, connect_taxops_readonly
from audit.normalizer import (
    keys_with_transposition,
    normalize_drake_client_name,
    normalize_entity,
    normalize_person,
)
from audit.util import dumps


def _load_other_office_keys(tax_log: Path) -> set[str]:
    wb = openpyxl.load_workbook(tax_log, read_only=True, data_only=True)
    name = None
    for s in wb.sheetnames:
        if s.strip().casefold() == "other office svcs".casefold():
            name = s
            break
    keys: set[str] = set()
    if not name:
        wb.close()
        return keys
    ws = wb[name]
    for row in ws.iter_rows(values_only=True):
        for v in row[:6]:
            if v is None:
                continue
            s = str(v).strip()
            if not s:
                continue
            for k in normalize_entity(s).keys:
                keys.add(k)
            # also person-style if comma
            if "," in s:
                n = normalize_drake_client_name(s)
                for a, b in n.match_keys:
                    keys.add(f"{a}|{b}")
    wb.close()
    return keys


def _load_drake_non1040_keys(drake: Path) -> set[str]:
    wb = openpyxl.load_workbook(drake, read_only=True, data_only=True)
    ws = wb.active
    keys: set[str] = set()
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        vals = list(row)
        name = str(vals[1] or "").strip() if len(vals) > 1 else ""
        rtype = str(vals[2] or "").strip().upper() if len(vals) > 2 else ""
        if not name or rtype in ("", "1040"):
            continue
        for k in normalize_entity(name).keys:
            keys.add(k)
        n = normalize_drake_client_name(name)
        for a, b in n.match_keys:
            keys.add(f"{a}|{b}")
    wb.close()
    return keys


def _load_log_other_year_keys(tax_log: Path) -> set[tuple[str, str]]:
    wb = openpyxl.load_workbook(tax_log, read_only=True, data_only=True)
    ws = wb[config.SHEET_INDIVIDUALS]
    keys: set[tuple[str, str]] = set()
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < config.LOG_DATA_START_ROW:
            continue
        vals = list(row)
        last = str(vals[config.LOG_LAST_COL] or "").strip() if len(vals) > config.LOG_LAST_COL else ""
        first = (
            str(vals[config.LOG_FIRST_COL] or "").strip()
            if len(vals) > config.LOG_FIRST_COL
            else ""
        )
        yr = vals[config.LOG_YR_COL] if len(vals) > config.LOG_YR_COL else None
        from audit.ingest import normalize_yr

        yn = normalize_yr(yr)
        if yn == 2025 or yn is None:
            continue
        if not last and not first:
            continue
        for k in keys_with_transposition(last, first):
            keys.add(k)
    wb.close()
    return keys


def run_m93(
    *,
    audit_db: Path,
    run_id: int,
    snapshot: Path,
    tax_log: Path,
    drake: Path,
) -> dict[str, Any]:
    """Build multi-label attrs for unmatched TY2025; write audit_gap_ty2025 refresh."""
    aconn = connect_audit(audit_db)
    tconn = connect_taxops_readonly(snapshot)
    try:
        # Unmatched = gap rows whose cause != matched_to_drake_or_log from prior run,
        # OR recompute: TY2025 returns whose client is not in matched set.
        matched_stage = set()
        for r in aconn.execute(
            """
            SELECT right_id FROM audit_match
            WHERE run_id=? AND pair IN ('DRAKE_TAXOPS','LOG_TAXOPS') AND right_kind='taxops'
            """,
            (run_id,),
        ):
            matched_stage.add(int(r[0]))
        stage_to_client = {
            int(r["id"]): int(r["client_id"])
            for r in aconn.execute(
                "SELECT id, client_id FROM stage_taxops_client WHERE run_id=?",
                (run_id,),
            )
        }
        matched_clients = {stage_to_client[s] for s in matched_stage if s in stage_to_client}

        # import batches (only 3)
        batches = list(
            tconn.execute(
                "SELECT id, filename, imported_at, file_hash FROM import_batches ORDER BY id"
            )
        )
        # Map return → batch via status_events.source_file or notes — best effort:
        # returns have no import_batch_id; use earliest status_events.source_file match to batch filename
        batch_by_return: dict[int, Optional[int]] = {}
        for r in tconn.execute(
            """
            SELECT return_id, source_file FROM status_events
            WHERE source_file IS NOT NULL AND source_file != ''
            ORDER BY id
            """
        ):
            rid = int(r["return_id"])
            if rid in batch_by_return:
                continue
            sf = str(r["source_file"])
            bid = None
            for b in batches:
                if b["filename"] and b["filename"] in sf:
                    bid = int(b["id"])
                    break
            batch_by_return[rid] = bid

        spouse_clients = {
            int(r[0])
            for r in aconn.execute(
                "SELECT DISTINCT client_id FROM stage_taxops_spouse WHERE run_id=?",
                (run_id,),
            )
        }
        other_year = {
            int(r[0])
            for r in aconn.execute(
                """
                SELECT DISTINCT client_id FROM stage_taxops_return
                WHERE run_id=? AND tax_year IS NOT NULL AND tax_year != 2025
                """,
                (run_id,),
            )
        }
        clients = {
            int(r["client_id"]): r
            for r in aconn.execute(
                "SELECT * FROM stage_taxops_client WHERE run_id=?", (run_id,)
            )
        }

        oos_keys = _load_other_office_keys(tax_log)
        non1040_keys = _load_drake_non1040_keys(drake)
        other_yr_log = _load_log_other_year_keys(tax_log)

        # Clear prior gap rows for this run and rewrite
        aconn.execute("DELETE FROM audit_gap_ty2025 WHERE run_id=?", (run_id,))

        rows_out = []
        ty = aconn.execute(
            "SELECT * FROM stage_taxops_return WHERE run_id=? AND tax_year=2025",
            (run_id,),
        ).fetchall()

        for r in ty:
            cid = int(r["client_id"])
            if cid in matched_clients:
                continue
            c = clients.get(cid)
            last = (c["last_name"] if c else "") or ""
            first = (c["first_name"] if c else "") or ""
            n = normalize_person(last, first)
            person_keys = set(n.match_keys)
            entity_keys = set(normalize_entity(last).keys) | set(
                normalize_entity(f"{last} {first}".strip()).keys
            )
            person_pipe = {f"{a}|{b}" for a, b in person_keys}

            has_log = 1 if r["log_number"] not in (None, "", "0", 0) else 0
            attrs = {
                "client_status": (r["client_status"] or "").strip() or None,
                "has_log_number": bool(has_log),
                "created_at": r["created_at"],
                "import_batch_id": batch_by_return.get(int(r["return_id"])),
                "has_other_year": cid in other_year,
                "has_spouses_row": cid in spouse_clients,
                "matches_other_office_svcs": bool(
                    entity_keys & oos_keys or person_pipe & oos_keys
                ),
                "matches_drake_non1040": bool(
                    entity_keys & non1040_keys or person_pipe & non1040_keys
                ),
                "matches_log_other_year": bool(person_keys & other_yr_log),
            }
            # combination key for cross-tab
            combo = (
                f"status={attrs['client_status'] or 'NULL'}"
                f"|log={attrs['has_log_number']}"
                f"|batch={attrs['import_batch_id']}"
                f"|other_yr={attrs['has_other_year']}"
                f"|spouse={attrs['has_spouses_row']}"
                f"|oos={attrs['matches_other_office_svcs']}"
                f"|non1040={attrs['matches_drake_non1040']}"
                f"|log_oy={attrs['matches_log_other_year']}"
            )
            rows_out.append(
                {
                    "return_id": int(r["return_id"]),
                    "client_id": cid,
                    "attrs": attrs,
                    "combo": combo,
                }
            )
            aconn.execute(
                """
                INSERT INTO audit_gap_ty2025 (
                  run_id, return_id, client_id, cause, client_status, has_log_number,
                  created_at, import_batch_id, has_other_year, other_office_plausible,
                  detail_json
                ) VALUES (?,?,?,?,?,?,?,?,?,?,?)
                """,
                (
                    run_id,
                    int(r["return_id"]),
                    cid,
                    combo,  # cause now holds combination label
                    attrs["client_status"],
                    has_log,
                    attrs["created_at"],
                    attrs["import_batch_id"],
                    1 if attrs["has_other_year"] else 0,
                    1 if attrs["matches_other_office_svcs"] else 0,
                    dumps(attrs),
                ),
            )

        aconn.commit()

        # created_at histogram by month
        month_hist = Counter()
        ts_counts = Counter()
        for row in rows_out:
            ca = row["attrs"]["created_at"] or ""
            month = str(ca)[:7] if ca else "NULL"
            month_hist[month] += 1
            if ca:
                ts_counts[str(ca)] += 1
        bulk_timestamps = {t: c for t, c in ts_counts.items() if c > 20}

        combo_counts = Counter(r["combo"] for r in rows_out)
        ranked = combo_counts.most_common()
        total = len(rows_out)
        covered = 0
        top = []
        for combo, cnt in ranked:
            top.append({"combo": combo, "count": cnt})
            covered += cnt
            if covered / total >= 0.80 if total else True:
                break

        # Named populations with hypotheses
        populations = []
        for item in top:
            combo = item["combo"]
            cnt = item["count"]
            # Parse flags
            parts = dict(p.split("=", 1) for p in combo.split("|"))
            hyp = []
            if parts.get("log") == "False":
                hyp.append("never received an office log number")
            if parts.get("batch") not in (None, "None"):
                hyp.append(f"tied to import_batch {parts.get('batch')}")
            else:
                hyp.append("no import_batch linkage via status_events.source_file")
            if parts.get("oos") == "True":
                hyp.append("name matches Other Office Svcs — likely non-prep work")
            if parts.get("non1040") == "True":
                hyp.append("name matches a Drake entity type — may be mis-grained as 1040 return")
            if parts.get("other_yr") == "True":
                hyp.append("client also has other-year returns — carry-forward / multi-year artifact")
            if parts.get("spouse") == "True":
                hyp.append("client has spouses row — real household, not a phantom stub")
            if parts.get("log_oy") == "True":
                hyp.append("name appears on Tax Log at a year other than 25")
            st = parts.get("status")
            if st and st not in ("NULL",):
                hyp.append(f"workflow status is {st}")
            populations.append(
                {
                    "name": combo,
                    "count": cnt,
                    "hypothesis": "; ".join(hyp) if hyp else "insufficient signal",
                }
            )

        return {
            "unmatched_total": total,
            "expected_227_note": "prior PHANTOM count was 227; this recompute may differ slightly with match set",
            "month_histogram": dict(sorted(month_hist.items())),
            "bulk_timestamps_gt_20": {k: v for k, v in list(bulk_timestamps.items())[:20]},
            "bulk_timestamp_count": len(bulk_timestamps),
            "top_combinations": top,
            "top_coverage": round(covered / total, 4) if total else 0,
            "populations": populations,
            "import_batches": [
                {"id": int(b["id"]), "filename": b["filename"], "imported_at": b["imported_at"]}
                for b in batches
            ],
            "attr_marginals": {
                "has_log_number": sum(1 for r in rows_out if r["attrs"]["has_log_number"]),
                "no_log_number": sum(1 for r in rows_out if not r["attrs"]["has_log_number"]),
                "has_other_year": sum(1 for r in rows_out if r["attrs"]["has_other_year"]),
                "has_spouses_row": sum(1 for r in rows_out if r["attrs"]["has_spouses_row"]),
                "matches_other_office_svcs": sum(
                    1 for r in rows_out if r["attrs"]["matches_other_office_svcs"]
                ),
                "matches_drake_non1040": sum(
                    1 for r in rows_out if r["attrs"]["matches_drake_non1040"]
                ),
                "matches_log_other_year": sum(
                    1 for r in rows_out if r["attrs"]["matches_log_other_year"]
                ),
                "by_status": dict(
                    Counter(r["attrs"]["client_status"] or "NULL" for r in rows_out)
                ),
                "by_batch": dict(
                    Counter(str(r["attrs"]["import_batch_id"]) for r in rows_out)
                ),
            },
        }
    finally:
        aconn.close()
        tconn.close()


if __name__ == "__main__":
    # Discover latest run
    import glob

    dbs = sorted(Path(r"T:\audit").glob("audit_*.sqlite"), key=lambda p: p.stat().st_mtime)
    db = dbs[-1]
    conn = connect_audit(db)
    run_id = conn.execute("SELECT MAX(id) FROM audit_run").fetchone()[0]
    conn.close()
    out = run_m93(
        audit_db=db,
        run_id=int(run_id),
        snapshot=Path(r"T:\audit\snapshots\taxops_snapshot_20260731.sqlite"),
        tax_log=Path(
            r"C:\Users\Windows 10\OneDrive - Xcel Financial Services LLC"
            r"\Shared\Logs\TAX LOG 2025 Live.xlsx"
        ),
        drake=Path(r"C:\Users\Windows 10\Desktop\CLIENTS.xlsx"),
    )
    Path(r"T:\audit\output\m93_gap_causal.json").write_text(
        json.dumps(out, indent=2, default=str), encoding="utf-8"
    )
    # console: no PII
    print(json.dumps(out, indent=2, default=str))
