"""M6 — findings workbook (openpyxl, /export house style)."""

from __future__ import annotations

import sqlite3
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from audit.db import connect_audit
from audit.util import stamp_day

HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
DATA_FONT = Font(name="Calibri", size=11)
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Border(bottom=Side(style="thin", color="E2E8F0"))

PASTEL = {
    "NEEDS_HUMAN": "FECACA",
    "DUPLICATE_CLIENT": "FDE68A",
    "MISSING_IN_TAXOPS": "FECACA",
    "PHANTOM_IN_TAXOPS": "E0E7FF",
    "PREPARED_NOT_LOGGED": "FEF9C3",
    "LOGGED_NOT_PREPARED": "CCFBF1",
    "SPOUSE_MISSING": "FCE7F3",
    "SPOUSE_STORE_DIVERGENCE": "E0F2FE",
    "NAME_TRUNCATED": "F1F5F9",
    "CONFIRMED_MATCH": "D1FAE5",
}


def _style_header(ws, cols):
    for i, (label, width) in enumerate(cols, 1):
        cell = ws.cell(1, i, label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 22


def write_workbook(audit_db: Path, run_id: int, out_dir: Path) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    out = out_dir / f"client_audit_findings_{stamp_day()}_run{run_id}.xlsx"

    conn = connect_audit(audit_db)
    try:
        gap = conn.execute(
            """
            SELECT cause, COUNT(*) c FROM audit_gap_ty2025
            WHERE run_id=? GROUP BY cause ORDER BY c DESC
            """,
            (run_id,),
        ).fetchall()
        findings = conn.execute(
            """
            SELECT * FROM audit_finding WHERE run_id=? ORDER BY severity ASC, id ASC
            """,
            (run_id,),
        ).fetchall()
        by_type: dict[str, list] = defaultdict(list)
        for f in findings:
            by_type[f["finding_type"]].append(f)

        wb = openpyxl.Workbook()
        # Summary first
        ws = wb.active
        ws.title = "Summary"
        _style_header(ws, [("Section", 28), ("Key", 40), ("Count", 12)])
        row = 2
        ws.cell(row, 1, "TY2025 gap by cause").font = Font(name="Calibri", bold=True, size=11)
        row = 3
        total = 0
        for g in gap:
            ws.cell(row, 1, "TY2025_GAP")
            ws.cell(row, 2, g["cause"])
            ws.cell(row, 3, g["c"])
            for col in range(1, 4):
                ws.cell(row, col).font = DATA_FONT
                ws.cell(row, col).border = THIN
            total += int(g["c"])
            row += 1
        ws.cell(row, 1, "TY2025_GAP")
        ws.cell(row, 2, "TOTAL")
        ws.cell(row, 3, total)
        row += 2
        ws.cell(row, 1, "Findings by type").font = Font(name="Calibri", bold=True, size=11)
        row += 1
        for ftype, rows_ in sorted(by_type.items(), key=lambda x: (x[1][0]["severity"], x[0])):
            ws.cell(row, 1, "FINDING")
            ws.cell(row, 2, ftype)
            ws.cell(row, 3, len(rows_))
            fill = PASTEL.get(ftype)
            if fill:
                ws.cell(row, 2).fill = PatternFill("solid", fgColor=fill)
            for col in range(1, 4):
                ws.cell(row, col).font = DATA_FONT
                ws.cell(row, col).border = THIN
            row += 1

        # One tab per finding type (Excel sheet name max 31)
        for ftype, rows_ in sorted(by_type.items(), key=lambda x: x[1][0]["severity"]):
            title = ftype[:31]
            wsf = wb.create_sheet(title)
            cols = [
                ("ID", 8), ("Subtype", 28), ("Severity", 10),
                ("Subject kind", 16), ("Subject id", 12), ("Tax year", 10),
                ("Detail JSON", 48), ("Source refs", 36),
            ]
            _style_header(wsf, cols)
            fill = PASTEL.get(ftype)
            for ri, f in enumerate(rows_, 2):
                vals = [
                    f["id"], f["subtype"], f["severity"], f["subject_kind"],
                    f["subject_id"], f["tax_year"], f["detail_json"], f["source_refs"],
                ]
                for ci, v in enumerate(vals, 1):
                    cell = wsf.cell(ri, ci, v)
                    cell.font = DATA_FONT
                    cell.border = THIN
                    if fill and ci == 2:
                        cell.fill = PatternFill("solid", fgColor=fill)

        wb.save(out)
        return out
    finally:
        conn.close()
