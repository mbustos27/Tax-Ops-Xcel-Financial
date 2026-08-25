"""A0 — explicit baseline lock, drift detection, TaxOps path resolution.

Never writes to TaxOps. Writes only under T:\\audit\\.
"""

from __future__ import annotations

import json
import os
import sqlite3
from dataclasses import asdict, dataclass, field
from datetime import datetime, timezone
from pathlib import Path
from typing import Any, Optional

import openpyxl

from audit import config
from audit.util import dumps, sha256_file, utc_now

# Persistent cross-run memory for drift (not TaxOps).
BASELINE_MEMORY_PATH = config.AUDIT_ROOT / "baseline_memory.json"
A0_REPORT_PATH = config.AUDIT_ROOT / "investigation" / "A0-baseline.md"

# Candidate CSM exports (investigation I0). Order is documentation only —
# selection is explicit, never silent fallback.
CSM_CANDIDATES = {
    "onedrive_ty2025": Path(
        r"C:\Users\Windows 10\OneDrive - Xcel Financial Services LLC\CLIENTS.xlsx"
    ),
    "desktop_ty2025": Path(r"C:\Users\Windows 10\Desktop\CLIENTS.xlsx"),
    "csvfiles_ty2024": Path(r"T:\taxops\CSVFILES\2024 CLIENTS.xlsx"),
}

# Prefill importer consumed TY2024 CSM (Aug 7) — lag vs audit TY2025 CSM.
PREFILL_CSM_PATH = CSM_CANDIDATES["csvfiles_ty2024"]

TAXOPS_PATH_CANDIDATES = (
    Path(r"T:\taxops\taxops.db"),
    Path(r"\\Xcel-server\taxops\taxops\taxops.db"),
    Path(r"C:\TaxOps\taxops\taxops.db"),  # server-local NSSM path; often missing on workstations
)


@dataclass
class SourceMeta:
    role: str
    path: str
    exists: bool
    size: Optional[int] = None
    mtime_utc: Optional[str] = None
    sha256: Optional[str] = None
    row_count: Optional[int] = None
    extra: dict[str, Any] = field(default_factory=dict)


@dataclass
class BaselineLock:
    locked_at: str
    authoritative_drake_key: str
    authoritative_drake_path: str
    authoritative_drake_reason: str
    tax_log_path: str
    authoritative_taxops_path: str
    taxops_path_resolution: str  # RESOLVED | AUTHORITATIVE_DB_UNRESOLVED
    taxops_path_notes: list[str]
    sources: list[SourceMeta]
    drift_vs_previous: dict[str, Any]
    findings: list[dict[str, Any]]
    csm_prefill_lag_note: str
    preflight_ok: bool
    # Amendment 1 — Drake invoice-bearing export
    invoice_export: dict[str, Any] = field(default_factory=dict)
    open_questions: list[str] = field(default_factory=list)


def _mtime_utc(path: Path) -> str:
    return datetime.fromtimestamp(path.stat().st_mtime, tz=timezone.utc).strftime(
        "%Y-%m-%dT%H:%M:%SZ"
    )


def count_drake_rows(path: Path) -> tuple[int, dict[str, int], Optional[str]]:
    """Return (named_row_count, type_counts, max_last_change_str)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    types: dict[str, int] = {}
    n = 0
    max_lc: Optional[str] = None
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        vals = list(row)
        if not any(v is not None and str(v).strip() for v in vals):
            continue
        name = str(vals[1] or "").strip() if len(vals) > 1 else ""
        if not name or name.upper().startswith("TOTAL"):
            continue
        n += 1
        t = str(vals[2] or "").strip().upper() if len(vals) > 2 else ""
        types[t] = types.get(t, 0) + 1
        if len(vals) > 7 and vals[7] is not None:
            lc = str(vals[7])
            if max_lc is None or lc > max_lc:
                max_lc = lc
    wb.close()
    return n, types, max_lc


def count_log_named(path: Path, sheet: str = config.SHEET_INDIVIDUALS) -> int:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    if sheet not in wb.sheetnames:
        wb.close()
        raise FileNotFoundError(f"Tax Log missing sheet {sheet!r}")
    ws = wb[sheet]
    named = 0
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < config.LOG_DATA_START_ROW:
            continue
        vals = list(row)
        last = (
            str(vals[config.LOG_LAST_COL] or "").strip()
            if len(vals) > config.LOG_LAST_COL
            else ""
        )
        first = (
            str(vals[config.LOG_FIRST_COL] or "").strip()
            if len(vals) > config.LOG_FIRST_COL
            else ""
        )
        if last or first:
            named += 1
    wb.close()
    return named


def _meta_file(role: str, path: Path, *, hash_it: bool = True, row_count: Optional[int] = None,
               extra: Optional[dict] = None) -> SourceMeta:
    if not path.exists():
        return SourceMeta(role=role, path=str(path), exists=False, extra=extra or {})
    return SourceMeta(
        role=role,
        path=str(path),
        exists=True,
        size=path.stat().st_size,
        mtime_utc=_mtime_utc(path),
        sha256=sha256_file(path) if hash_it else None,
        row_count=row_count,
        extra=extra or {},
    )


def resolve_taxops_path() -> tuple[str, str, list[str], list[SourceMeta]]:
    """
    Returns (authoritative_path, resolution_code, notes, metas).

    Workstation: T:\\ and \\\\Xcel-server\\taxops\\... are the same file (samefile).
    C:\\TaxOps\\taxops\\taxops.db is the NSSM path on the server — often absent here.
    Whether server C:\\ equals the share cannot be proven from the workstation alone.
    """
    notes: list[str] = []
    metas: list[SourceMeta] = []
    existing: list[Path] = []
    for p in TAXOPS_PATH_CANDIDATES:
        m = _meta_file(f"taxops_candidate:{p}", p, hash_it=False)
        # Don't store full path as role with weird chars — simplify
        m.role = "taxops_candidate"
        m.extra = {"label": str(p)}
        metas.append(m)
        if p.exists():
            existing.append(p)

    t_path = Path(r"T:\taxops\taxops.db")
    unc = Path(r"\\Xcel-server\taxops\taxops\taxops.db")
    c_path = Path(r"C:\TaxOps\taxops\taxops.db")

    if t_path.exists() and unc.exists():
        try:
            same = os.path.samefile(t_path, unc)
        except OSError:
            same = False
        notes.append(
            f"T:\\taxops\\taxops.db samefile UNC: {same} "
            f"(size={t_path.stat().st_size}, mtime_utc={_mtime_utc(t_path)})"
        )
        if same:
            notes.append(
                "Share mapping confirmed: workstation T: and \\\\Xcel-server\\taxops\\taxops\\taxops.db "
                "are the same inode."
            )
        auth = str(unc)  # prefer UNC as canonical share path
    elif t_path.exists():
        auth = str(t_path)
        notes.append("UNC path not reachable; using T:\\taxops\\taxops.db")
    elif unc.exists():
        auth = str(unc)
        notes.append("T: not reachable; using UNC")
    else:
        auth = str(t_path)
        notes.append("No TaxOps DB reachable on T: or UNC")
        return auth, "AUTHORITATIVE_DB_UNRESOLVED", notes, metas

    if not c_path.exists():
        notes.append(
            "C:\\TaxOps\\taxops\\taxops.db does not exist on this workstation "
            "(expected: NSSM DB_PATH on the app server only)."
        )
        notes.append(
            "Cannot prove from this workstation whether server-local C:\\TaxOps\\... "
            "is the same file as the share — emitting AUTHORITATIVE_DB_UNRESOLVED for that pairing."
        )
        return auth, "AUTHORITATIVE_DB_UNRESOLVED", notes, metas

    # Rare: C: exists on this machine — try samefile against share
    share = Path(auth)
    try:
        same_c = os.path.samefile(c_path, share)
    except OSError:
        same_c = False
    notes.append(f"C:\\TaxOps samefile share: {same_c}")
    if same_c:
        return auth, "RESOLVED", notes, metas
    notes.append("C:\\TaxOps exists but is NOT the same file as the share.")
    return auth, "AUTHORITATIVE_DB_UNRESOLVED", notes, metas


def choose_authoritative_drake(
    *,
    cli_drake: Optional[Path] = None,
) -> tuple[str, Path, str, list[dict[str, Any]]]:
    """
    Explicit CSM baseline decision.

    Preference (from I0/I4 + task1_export_direction.json):
      OneDrive TY2025 (later by Drake Last Change, historically 1159)
      > Desktop TY2025 (1155, older)
      NEVER silently use TY2024 CSVFILES (wrong season for this three-way).

    If --drake is passed, that path wins only when it matches a known candidate
    or the operator path is recorded as an explicit override finding.
    """
    findings: list[dict[str, Any]] = []
    onedrive = CSM_CANDIDATES["onedrive_ty2025"]
    desktop = CSM_CANDIDATES["desktop_ty2025"]
    ty2024 = CSM_CANDIDATES["csvfiles_ty2024"]

    if cli_drake is not None:
        resolved = cli_drake.resolve() if cli_drake.exists() else cli_drake
        # Detect accidental TY2024 selection
        if ty2024.exists() and resolved.exists():
            try:
                if os.path.samefile(resolved, ty2024):
                    findings.append(
                        {
                            "type": "BASELINE_REJECTED_TY2024_CSM",
                            "detail": "CLI --drake pointed at TY2024 CSVFILES export; refusing as TY2025 three-way baseline.",
                        }
                    )
                    raise ValueError(
                        "Refusing TY2024 CLIENTS.xlsx as authoritative Drake baseline for TY2025 audit. "
                        "Pass OneDrive or Desktop TY2025 CSM."
                    )
            except OSError:
                pass
        if not resolved.exists():
            raise FileNotFoundError(f"--drake not found: {cli_drake}")
        key = "cli_override"
        for k, p in CSM_CANDIDATES.items():
            if p.exists():
                try:
                    if os.path.samefile(resolved, p):
                        key = k
                        break
                except OSError:
                    continue
        reason = (
            f"Explicit --drake {resolved}. "
            + (
                "Matches known OneDrive TY2025 candidate (preferred)."
                if key == "onedrive_ty2025"
                else "Matches known Desktop TY2025 candidate."
                if key == "desktop_ty2025"
                else "Operator override path (not one of the three I0 candidates)."
            )
        )
        if key == "cli_override":
            findings.append(
                {
                    "type": "BASELINE_CLI_OVERRIDE",
                    "detail": str(resolved),
                }
            )
        return key, resolved, reason, findings

    # No CLI: prefer OneDrive if present (investigation: later Last Change, EXPECTED 1159)
    if onedrive.exists():
        reason = (
            "Authoritative CSM = OneDrive CLIENTS.xlsx. "
            "Investigation task1_export_direction.json: later by Drake Last Change than Desktop; "
            "historically 1159 rows. Present on disk at firm OneDrive root."
        )
        return "onedrive_ty2025", onedrive, reason, findings

    findings.append(
        {
            "type": "BASELINE_PREFERRED_CSM_MISSING",
            "detail": str(onedrive),
        }
    )
    if desktop.exists():
        reason = (
            "FALLBACK: OneDrive TY2025 CSM not found; locking Desktop CLIENTS.xlsx (1155). "
            "This is older by Last Change than the OneDrive export used in Jul31 audit lock."
        )
        findings.append(
            {
                "type": "BASELINE_FALLBACK_DESKTOP",
                "detail": str(desktop),
            }
        )
        return "desktop_ty2025", desktop, reason, findings

    raise FileNotFoundError(
        "No TY2025 CSM export found (OneDrive or Desktop). "
        f"Checked: {onedrive}, {desktop}. TY2024 at {ty2024} is not eligible."
    )


def load_baseline_memory() -> dict[str, Any]:
    if not BASELINE_MEMORY_PATH.exists():
        return {}
    return json.loads(BASELINE_MEMORY_PATH.read_text(encoding="utf-8"))


def save_baseline_memory(payload: dict[str, Any]) -> None:
    BASELINE_MEMORY_PATH.write_text(dumps(payload), encoding="utf-8")


def compute_drift(current: dict[str, Any], previous: dict[str, Any]) -> tuple[dict[str, Any], list[dict]]:
    """Compare current observed counts to previous run. Drift → BASELINE_DRIFT finding, not crash."""
    findings: list[dict[str, Any]] = []
    drift: dict[str, Any] = {"previous_run_at": previous.get("locked_at"), "changes": {}}
    if not previous:
        drift["status"] = "NO_PREVIOUS"
        return drift, findings

    keys = (
        "drake_rows",
        "log_named_rows",
        "taxops_clients",
        "taxops_returns",
        "taxops_ty2025",
        "taxops_spouses",
    )
    prev_counts = previous.get("counts") or {}
    cur_counts = current
    for k in keys:
        pv, cv = prev_counts.get(k), cur_counts.get(k)
        if pv is None or cv is None:
            continue
        if pv != cv:
            drift["changes"][k] = {"previous": pv, "current": cv, "delta": cv - pv}
    if drift["changes"]:
        drift["status"] = "DRIFT"
        findings.append(
            {
                "type": "BASELINE_DRIFT",
                "detail": drift["changes"],
            }
        )
    else:
        drift["status"] = "STABLE"
    return drift, findings


def build_baseline_lock(
    *,
    tax_log_path: Path,
    cli_drake: Optional[Path] = None,
    taxops_snapshot: Optional[Path] = None,
) -> BaselineLock:
    findings: list[dict[str, Any]] = []
    sources: list[SourceMeta] = []

    key, drake_path, reason, choose_findings = choose_authoritative_drake(cli_drake=cli_drake)
    findings.extend(choose_findings)

    # Catalog all CSM candidates (hash + count)
    for cand_key, cand_path in CSM_CANDIDATES.items():
        if cand_path.exists():
            n, types, max_lc = count_drake_rows(cand_path)
            sources.append(
                _meta_file(
                    f"csm:{cand_key}",
                    cand_path,
                    row_count=n,
                    extra={"types": types, "max_last_change": max_lc, "candidate_key": cand_key},
                )
            )
        else:
            sources.append(_meta_file(f"csm:{cand_key}", cand_path, hash_it=False))

    d_count, d_types, d_max_lc = count_drake_rows(drake_path)
    # Ensure authoritative meta is present with role marker
    sources.append(
        _meta_file(
            "csm:AUTHORITATIVE",
            drake_path,
            row_count=d_count,
            extra={
                "types": d_types,
                "max_last_change": d_max_lc,
                "candidate_key": key,
                "reason": reason,
            },
        )
    )

    if not tax_log_path.exists():
        raise FileNotFoundError(f"Tax Log not found: {tax_log_path}")
    log_named = count_log_named(tax_log_path)
    sources.append(
        _meta_file("tax_log", tax_log_path, row_count=log_named, extra={"sheet": config.SHEET_INDIVIDUALS})
    )

    auth_taxops, resolution, taxops_notes, taxops_metas = resolve_taxops_path()
    sources.extend(taxops_metas)
    if resolution == "AUTHORITATIVE_DB_UNRESOLVED":
        findings.append(
            {
                "type": "AUTHORITATIVE_DB_UNRESOLVED",
                "detail": {
                    "authoritative_path_for_audit": auth_taxops,
                    "notes": taxops_notes,
                },
            }
        )

    # TaxOps live counts (read-only) + optional snapshot meta
    taxops_counts: dict[str, int] = {}
    page_count = None
    auth_p = Path(auth_taxops)
    if auth_p.exists():
        from audit.db import connect_taxops_readonly

        conn = connect_taxops_readonly(auth_p)
        try:
            taxops_counts = {
                "taxops_clients": int(conn.execute("SELECT COUNT(*) FROM clients").fetchone()[0]),
                "taxops_returns": int(conn.execute("SELECT COUNT(*) FROM returns").fetchone()[0]),
                "taxops_ty2025": int(
                    conn.execute("SELECT COUNT(*) FROM returns WHERE tax_year=2025").fetchone()[0]
                ),
                "taxops_spouses": int(conn.execute("SELECT COUNT(*) FROM spouses").fetchone()[0]),
            }
            page_count = int(conn.execute("PRAGMA page_count").fetchone()[0])
        finally:
            conn.close()
        sources.append(
            _meta_file(
                "taxops:AUTHORITATIVE",
                auth_p,
                hash_it=True,
                row_count=taxops_counts.get("taxops_clients"),
                extra={"page_count": page_count, **taxops_counts},
            )
        )

    if taxops_snapshot and taxops_snapshot.exists():
        snap_conn_meta = _meta_file(
            "taxops:snapshot",
            taxops_snapshot,
            hash_it=True,
            extra={"page_count": None},
        )
        try:
            from audit.db import connect_taxops_readonly

            sc = connect_taxops_readonly(taxops_snapshot)
            try:
                snap_conn_meta.extra["page_count"] = int(sc.execute("PRAGMA page_count").fetchone()[0])
            finally:
                sc.close()
        except sqlite3.Error:
            pass
        sources.append(snap_conn_meta)

    # Prefill lag note
    prefill_note = (
        "Prefill importer consumed T:\\taxops\\CSVFILES\\2024 CLIENTS.xlsx "
        f"(mtime {_mtime_utc(PREFILL_CSM_PATH) if PREFILL_CSM_PATH.exists() else 'MISSING'}, "
        "TY2024 universe). "
        f"Audit authoritative CSM is {key} at {drake_path} "
        f"(mtime {_mtime_utc(drake_path)}, TY2025). "
        "Three-way census vs prefill-linked TaxOps rows will mix export-lag with true data error — "
        "treat CSM↔prefill disagreements as EXPORT_LAG until seasons align. "
        "Amendment 1: prefer TAXPAYER.csv (separate first/last, Invoice Number) over CSM joint names "
        "for linkage identity when both exist."
    )
    if PREFILL_CSM_PATH.exists():
        sources.append(
            _meta_file(
                "csm:prefill_consumed_ty2024",
                PREFILL_CSM_PATH,
                row_count=next(
                    (s.row_count for s in sources if s.role == "csm:csvfiles_ty2024"),
                    None,
                ),
                extra={"note": "Not authoritative for TY2025 three-way"},
            )
        )

    # ── Amendment 1: TAXPAYER.csv invoice export ─────────────────────────
    from audit.invoice_export import (
        DEFAULT_TAXPAYER_INVOICE_PATH,
        SUPERSEDED_SPOUSE_EXPORTS,
        parse_taxpayer_invoice_csv,
        report_to_jsonable,
    )

    inv_path = DEFAULT_TAXPAYER_INVOICE_PATH
    inv_report = parse_taxpayer_invoice_csv(inv_path)
    invoice_blob = report_to_jsonable(inv_report)
    # Strip heavy findings from memory blob; keep summary counts
    findings.extend(
        f
        for f in inv_report.findings
        if f.get("type")
        in (
            "MALFORMED_LOG_NUMBER",
            "LOG_NUMBER_COLLISION",
            "PROFORMA_STALE_LOG",
            "INVOICE_EXPORT_MISSING",
            "INVOICE_HEADER_MISMATCH",
            "UNEXPECTED_EMPTY_INVOICE_ON_FULL_WIDTH",
        )
        # RAGGED_EXPORT_ROW is numerous — summarize once
    )
    ragged_n = sum(1 for f in inv_report.findings if f.get("type") == "RAGGED_EXPORT_ROW")
    if ragged_n:
        findings.append(
            {
                "type": "RAGGED_EXPORT_ROW",
                "detail": {
                    "n": ragged_n,
                    "histogram": inv_report.field_count_histogram,
                    "note": "Short lines padded for parse. On spouse_11, Invoice is last col "
                    "so shortfall drops the key; on link_5, Invoice is col 3.",
                },
            }
        )

    sources.append(
        _meta_file(
            "drake:TAXPAYER_invoice",
            inv_path,
            hash_it=False,  # already hashed inside parse
            row_count=inv_report.data_row_count,
            extra={
                "sha256": inv_report.sha256,
                "title_lines": inv_report.title_lines,
                "tax_year": inv_report.tax_year,
                "layout": inv_report.layout,
                "expected_col_count": inv_report.expected_col_count,
                "full_width": inv_report.full_width_count,
                "ragged": inv_report.ragged_count,
                "empty_invoice_full_width": inv_report.empty_invoice_full_width,
                "l0_ok_invoices": inv_report.valid_l0_invoice_count,
                "l0_ok_bare": inv_report.valid_l0_bare_count,
                "legacy_l0_invoice_count": inv_report.legacy_l0_invoice_count,
                "bare_log_max": inv_report.bare_log_max,
                "entity_full_with_invoice": inv_report.entity_full_with_invoice,
                "entity_blank_first": inv_report.entity_blank_first,
            },
        )
    )
    # Repair SourceMeta sha/mtime from report
    for s in sources:
        if s.role == "drake:TAXPAYER_invoice":
            s.sha256 = inv_report.sha256
            s.mtime_utc = inv_report.mtime_utc
            s.size = inv_report.size
            s.exists = inv_report.exists

    for old in SUPERSEDED_SPOUSE_EXPORTS:
        if old.exists():
            # Detect 10-col (no Invoice Number)
            try:
                head = old.read_text(encoding="utf-8-sig", errors="replace").splitlines()[:5]
                hdr = next((ln for ln in head if "Taxpayer" in ln and "," in ln), "")
                if "Invoice Number" not in hdr:
                    findings.append(
                        {
                            "type": "SUPERSEDED_EXPORT",
                            "detail": {
                                "path": str(old),
                                "superseded_by": str(inv_path),
                                "reason": "10-column spouse export without Invoice Number",
                            },
                        }
                    )
                    sources.append(
                        _meta_file(
                            "drake:spouse_export_SUPERSEDED",
                            old,
                            hash_it=True,
                            extra={"superseded_by": str(inv_path)},
                        )
                    )
            except OSError:
                pass

    open_questions = [
        "Can the Drake report writer emit Invoice Number alongside Status in one export? "
        "If yes, CSM stops being load-bearing for linkage (key + clean names + workflow status). "
        "Record answer on audit_run; do not assume.",
    ]

    counts_now = {
        "drake_rows": d_count,
        "log_named_rows": log_named,
        "invoice_full_width_rows": inv_report.full_width_count,
        "invoice_l0_ok": inv_report.valid_l0_invoice_count,
        **taxops_counts,
    }
    previous = load_baseline_memory()
    drift, drift_findings = compute_drift(counts_now, previous)
    findings.extend(drift_findings)

    lock = BaselineLock(
        locked_at=utc_now(),
        authoritative_drake_key=key,
        authoritative_drake_path=str(drake_path),
        authoritative_drake_reason=reason,
        tax_log_path=str(tax_log_path),
        authoritative_taxops_path=auth_taxops,
        taxops_path_resolution=resolution,
        taxops_path_notes=taxops_notes,
        sources=sources,
        drift_vs_previous=drift,
        findings=findings,
        csm_prefill_lag_note=prefill_note,
        preflight_ok=True,
        invoice_export=invoice_blob,
        open_questions=open_questions,
    )

    # Persist memory for next run's drift detection
    save_baseline_memory(
        {
            "locked_at": lock.locked_at,
            "authoritative_drake_key": key,
            "authoritative_drake_path": str(drake_path),
            "authoritative_drake_sha256": next(
                (s.sha256 for s in sources if s.role == "csm:AUTHORITATIVE"), None
            ),
            "tax_log_path": str(tax_log_path),
            "authoritative_taxops_path": auth_taxops,
            "taxpayer_invoice_path": str(inv_path),
            "taxpayer_invoice_sha256": inv_report.sha256,
            "counts": counts_now,
            "drake_types": d_types,
        }
    )
    return lock


def lock_to_dict(lock: BaselineLock) -> dict[str, Any]:
    d = asdict(lock)
    return d


def write_a0_report(lock: BaselineLock, dest: Path = A0_REPORT_PATH) -> Path:
    dest.parent.mkdir(parents=True, exist_ok=True)
    auth = next(s for s in lock.sources if s.role == "csm:AUTHORITATIVE")
    log = next(s for s in lock.sources if s.role == "tax_log")
    lines: list[str] = []
    A = lines.append
    A("# A0 — Baseline lock")
    A("")
    A(f"_Locked at: {lock.locked_at}_")
    A("")
    A("## Authoritative CSM (explicit decision)")
    A("")
    A(f"- **Key:** `{lock.authoritative_drake_key}`")
    A(f"- **Path:** `{lock.authoritative_drake_path}`")
    A(f"- **Reason:** {lock.authoritative_drake_reason}")
    A(f"- **Rows:** {auth.row_count}")
    A(f"- **SHA-256:** `{auth.sha256}`")
    A(f"- **mtime (UTC):** {auth.mtime_utc}")
    A(f"- **Max Last Change (cell):** {auth.extra.get('max_last_change')}")
    A(f"- **Types:** `{dumps(auth.extra.get('types'))}`")
    A("")
    A("### Candidates surveyed (no silent fallback)")
    A("")
    A("| Key | Exists | Rows | SHA-256 (prefix) | mtime UTC |")
    A("|---|---|---:|---|---|")
    for s in lock.sources:
        if s.role.startswith("csm:") and s.role not in (
            "csm:AUTHORITATIVE",
            "csm:prefill_consumed_ty2024",
        ):
            sha = ((s.sha256 or "")[:16] + "…") if s.sha256 else "—"
            A(
                f"| `{s.extra.get('candidate_key', s.role)}` | {s.exists} | "
                f"{s.row_count if s.row_count is not None else '—'} | `{sha}` | "
                f"{s.mtime_utc or '—'} |"
            )
    A("")
    A("**TY2024 `CSVFILES/2024 CLIENTS.xlsx` is not eligible** as the TY2025 three-way baseline.")
    A("")
    A("## Tax Log")
    A("")
    A(f"- **Path:** `{lock.tax_log_path}`")
    A(f"- **Named rows (XCEL 2025):** {log.row_count}")
    A(f"- **SHA-256:** `{log.sha256}`")
    A(f"- **mtime (UTC):** {log.mtime_utc}")
    A("")
    A("## TaxOps DB path resolution")
    A("")
    A(f"- **Authoritative path for audit reads:** `{lock.authoritative_taxops_path}`")
    A(f"- **Resolution code:** `{lock.taxops_path_resolution}`")
    for n in lock.taxops_path_notes:
        A(f"- {n}")
    tax = next((s for s in lock.sources if s.role == "taxops:AUTHORITATIVE"), None)
    if tax:
        A(f"- **SHA-256:** `{tax.sha256}`")
        A(f"- **mtime (UTC):** {tax.mtime_utc}")
        A(f"- **page_count:** {(tax.extra or {}).get('page_count')}")
        A(f"- **clients/returns:** {(tax.extra or {}).get('taxops_clients')} / {(tax.extra or {}).get('taxops_returns')}")
    A("")
    A("## CSM ↔ prefill export lag")
    A("")
    A(lock.csm_prefill_lag_note)
    A("")
    A("## Amendment 1 — Drake TAXPAYER.csv (Invoice Number)")
    A("")
    inv = lock.invoice_export or {}
    if not inv:
        A("_Invoice export not registered on this lock._")
    else:
        A(f"- **Path:** `{inv.get('path')}` (canonical: `T:\\audit\\investigation\\TAXPAYER.csv`)")
        A(f"- **Exists:** {inv.get('exists')}")
        A(f"- **SHA-256:** `{inv.get('sha256')}`")
        A(f"- **mtime (UTC):** {inv.get('mtime_utc')}")
        A(f"- **Title lines:** `{inv.get('title_lines')}`")
        A(f"- **Tax year / season prefix:** {inv.get('tax_year')} / `{inv.get('season_prefix')}`")
        layout = inv.get("layout") or "spouse_11"
        exp_cols = inv.get("expected_col_count") or (5 if layout == "link_5" else 11)
        A(f"- **Layout:** `{layout}` (expected {exp_cols} columns)")
        A(f"- **Data rows:** {inv.get('data_row_count')}")
        A(f"- **Full-width ({exp_cols} fields):** {inv.get('full_width_count')} "
          f"({inv.get('invoice_coverage_full_width_pct')}% of data rows)")
        if layout == "link_5":
            A(f"- **Explicit blank Invoice (full-width):** {inv.get('empty_invoice_full_width')} "
              f"— not ragged; no key on that row")
            A(f"- **Ragged (short) rows:** {inv.get('ragged_count')} — Invoice is col 3; "
              f"short lines still lose trailing fields, not the invoice key")
        else:
            A(f"- **Ragged (short) rows:** {inv.get('ragged_count')} — shortfall is **export defect**, "
              f"not missing office invoices (Invoice Number is last column; short lines drop the key)")
        A(f"- **Field-count histogram:** `{dumps(inv.get('field_count_histogram'))}`")
        A(f"- **Name lengths:** first max={inv.get('first_name_max_len')} "
          f"(at39={inv.get('first_at_39')}, at40={inv.get('first_at_40')}); "
          f"last max={inv.get('last_name_max_len')} "
          f"(at39={inv.get('last_at_39')}, at40={inv.get('last_at_40')}) — "
          f"**no 40-char CSM ceiling in this export**")
        A("")
        A("### Entity-specific invoice coverage")
        A("")
        A(f"- Blank first name (entities): **{inv.get('entity_blank_first')}**")
        A(f"- Of those, ragged (key lost): **{inv.get('entity_ragged')}**")
        A(f"- Entities with readable invoice (full-width): **{inv.get('entity_full_with_invoice')}**")
        A("")
        A("### Invoice format (Amendment 2 C1/C2 — bare-log eligibility)")
        A("")
        A(f"- Length histogram (full-width only): `{dumps(inv.get('invoice_len_histogram'))}`")
        A(f"- Observed bare-log ceiling: **{inv.get('bare_log_max')}** "
          f"`{dumps(inv.get('bare_log_ceiling_meta') or {})}`")
        A(f"- **L0 before (legacy `^\\d{{6}}$`):** {inv.get('legacy_l0_invoice_count')}")
        A(f"- **L0 after (bare-log range, not collision):** raw invoices="
          f"**{inv.get('valid_l0_invoice_count')}**, distinct bare="
          f"**{inv.get('valid_l0_bare_count')}**")
        A(f"- Malformed (empty/out-of-range bare after normalize): "
          f"**{len(inv.get('malformed_invoices') or [])}** shown (excluded from L0)")
        for m in (inv.get("malformed_invoices") or [])[:12]:
            A(
                f"  - raw=`{m.get('invoice')}` bare=`{m.get('bare_log')}` "
                f"len={m.get('length')} reason={m.get('reason')} — {m.get('name')}"
            )
        A(f"- Collisions (same **bare log** → >1 distinct taxpayer): "
          f"**{len(inv.get('collisions') or [])}**")
        for c in (inv.get("collisions") or [])[:10]:
            names = c.get("claimant_names") or [
                (x.get("name") if isinstance(x, dict) else x) for x in (c.get("claimants") or [])
            ]
            A(
                f"  - bare=`{c.get('bare_log')}` raws=`{c.get('raw_invoices')}` "
                f"n={c.get('n_taxpayers')}: {names}"
            )
        A(f"- Proforma stale list (superseded by bare-range gate): "
          f"**{len(inv.get('proforma_stale') or [])}**")
    A("")
    A("## Drift vs previous run")
    A("")
    A(f"- **Status:** `{lock.drift_vs_previous.get('status')}`")
    A(f"- **Previous lock:** {lock.drift_vs_previous.get('previous_run_at')}")
    if lock.drift_vs_previous.get("changes"):
        A("")
        A("| Metric | Previous | Current | Delta |")
        A("|---|---:|---:|---:|")
        for k, v in lock.drift_vs_previous["changes"].items():
            A(f"| {k} | {v['previous']} | {v['current']} | {v['delta']} |")
        A("")
        A("Drift emits finding type `BASELINE_DRIFT` — **does not crash** the run.")
    else:
        A("")
        A("No count changes vs previous memory (or no previous memory).")
    A("")
    A("## Findings emitted")
    A("")
    # Summarize high-volume types
    from collections import Counter
    fc = Counter(f.get("type") for f in lock.findings)
    A(f"Counts by type: `{dumps(dict(fc))}`")
    A("")
    shown = 0
    for f in lock.findings:
        if f.get("type") == "RAGGED_EXPORT_ROW" and isinstance(f.get("detail"), dict) and "n" in (f.get("detail") or {}):
            A(f"- **{f['type']}:** `{dumps(f.get('detail'))}`")
            shown += 1
        elif f.get("type") != "RAGGED_EXPORT_ROW":
            if shown < 40:
                A(f"- **{f['type']}:** `{dumps(f.get('detail'))}`")
                shown += 1
    A("")
    A("## Open questions (carry to audit_run)")
    A("")
    for q in lock.open_questions or []:
        A(f"- {q}")
    A("")
    A("## Hard-constant retirement")
    A("")
    A(
        "`EXPECTED_DRAKE_ROWS` / `EXPECTED_LOG_NAMED_ROWS` / TaxOps expected counts in "
        "`audit/config.py` are no longer preflight crash gates. They remain as "
        "*historical reference only* (`HISTORICAL_*`); live comparison is previous-run memory "
        f"(`{BASELINE_MEMORY_PATH.name}`) → `BASELINE_DRIFT`."
    )
    A("")
    A("## Loud correction to I4")
    A("")
    A(
        "I4 stated no shared key exists between Drake and TaxOps. **Superseded by Amendment 1:** "
        "Drake Invoice Number holds the Tax Log number and is exportable via TAXPAYER.csv. "
        "L0 is now three-way on `(invoice_number, tax_year)`."
    )
    A("")
    A("## Files touched (A0 + Amendment 1)")
    A("")
    A("- `audit/baseline.py`")
    A("- `audit/invoice_export.py` (new)")
    A("- `audit/db.py`, `audit/config.py`, `audit/preflight.py`, `audit/__main__.py`")
    A("- `audit/baseline_memory.json`")
    A("- `audit/investigation/TAXPAYER.csv` (canonical invoice export copy)")
    A("- `audit/investigation/A0-baseline.md` (this file)")
    A("")
    dest.write_text("\n".join(lines), encoding="utf-8")
    return dest
