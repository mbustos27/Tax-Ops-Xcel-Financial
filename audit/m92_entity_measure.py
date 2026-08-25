"""M9.2 — measure entity match rate before/after normalize_entity."""

from __future__ import annotations

import json
from collections import Counter
from pathlib import Path

import openpyxl

from audit import config
from audit.match import Side, match_entities, drake_log_match_rate
from audit.normalizer import normalize_entity


def _resolve_sheet(wb, wanted: str):
    w = wanted.strip().casefold()
    for name in wb.sheetnames:
        if name.strip().casefold() == w:
            return name
    return None


def load_drake_entities(path: Path) -> list[Side]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    out = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        vals = list(row)
        name = str(vals[1] or "").strip() if len(vals) > 1 else ""
        if not name or name.upper().startswith("TOTAL"):
            continue
        rtype = str(vals[2] or "").strip().upper() if len(vals) > 2 else ""
        if rtype not in config.ENTITY_DRAKE_TYPES:
            continue
        id_raw = str(vals[0] or "")
        digits = "".join(c for c in id_raw if c.isdigit())
        last4 = digits[-4:] if len(digits) >= 4 else None
        out.append(
            Side(
                kind="drake",
                id=i,
                last=name,
                first="",
                display=name,
                last4=last4,
                is_entity=True,
            )
        )
    wb.close()
    return out


CORE_ENTITY_SHEETS = ("1120 CORP LIST", "1120 S LIST", "1065 & LLC LIST")
EXT_ENTITY_SHEETS = ("EXT 1120", "EXT 1120S, 1065'S")


def load_log_entities(path: Path) -> tuple[list[Side], dict]:
    """Core business lists + XCEL blank-FIRST + EXT sheets (tagged)."""
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    out: list[Side] = []
    stats = {
        "core_rows": 0,
        "ext_rows": 0,
        "xcel_blank_first": 0,
        "pool_tags": [],
    }
    nid = 1

    def add_sheet(wanted: str, tag: str) -> int:
        nonlocal nid
        resolved = _resolve_sheet(wb, wanted)
        if not resolved:
            return 0
        ws = wb[resolved]
        n = 0
        for i, row in enumerate(ws.iter_rows(values_only=True), 1):
            vals = list(row)
            if not any(v is not None and str(v).strip() for v in vals):
                continue
            name_blob = ""
            for v in vals[:8]:
                if v is None:
                    continue
                s = str(v).strip()
                if not s or s.upper() in ("NAME", "EIN", "CLIENT", "TYPE", "STATUS"):
                    continue
                if s.replace(".", "").isdigit():
                    continue
                name_blob = s
                break
            if not name_blob:
                continue
            n += 1
            out.append(
                Side(
                    kind="log",
                    id=nid,
                    last=name_blob,
                    first="",
                    display=name_blob,
                    is_entity=True,
                )
            )
            stats["pool_tags"].append(tag)
            nid += 1
        return n

    for sn in CORE_ENTITY_SHEETS:
        stats["core_rows"] += add_sheet(sn, "core")
    for sn in EXT_ENTITY_SHEETS:
        stats["ext_rows"] += add_sheet(sn, "ext")

    sheet = _resolve_sheet(wb, config.SHEET_INDIVIDUALS)
    ws = wb[sheet]
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < config.LOG_DATA_START_ROW:
            continue
        vals = list(row)
        last = str(vals[config.LOG_LAST_COL] or "").strip() if len(vals) > config.LOG_LAST_COL else ""
        first = (
            str(vals[config.LOG_FIRST_COL] or "").strip()
            if len(vals) > config.LOG_FIRST_COL
            else ""
        )
        if last and not first:
            stats["xcel_blank_first"] += 1
            out.append(
                Side(
                    kind="log",
                    id=nid,
                    last=last,
                    first="",
                    display=last,
                    is_entity=True,
                )
            )
            stats["pool_tags"].append("xcel_blank_first")
            nid += 1
    wb.close()
    return out, stats


def characterize_residual(lefts: list[Side], rights: list[Side]) -> dict:
    from audit.normalizer import normalize_entity

    # Build indexes
    suf_idx: dict[str, int] = Counter()
    strip_idx: dict[str, int] = Counter()
    for r in rights:
        ne = normalize_entity(r.display or r.last)
        for k in ne.keys_with_suffix:
            suf_idx[k] += 1
        for k in ne.keys:
            strip_idx[k] += 1
    causes = Counter()
    for left in lefts:
        ne = normalize_entity(left.display or left.last)
        if not ne.keys and not ne.keys_with_suffix:
            causes["empty_after_normalize"] += 1
            continue
        suf_hits = sum(suf_idx.get(k, 0) for k in ne.keys_with_suffix)
        strip_hits = sum(strip_idx.get(k, 0) for k in ne.keys)
        if suf_hits == 0 and strip_hits == 0:
            causes["no_key_overlap"] += 1
        elif suf_hits > 1 or strip_hits > 1:
            causes["ambiguous_multi_candidate"] += 1
        else:
            causes["unresolved_other"] += 1
    return dict(causes)


def main():
    drake = Path(r"C:\Users\Windows 10\Desktop\CLIENTS.xlsx")
    log = Path(
        r"C:\Users\Windows 10\OneDrive - Xcel Financial Services LLC"
        r"\Shared\Logs\TAX LOG 2025 Live.xlsx"
    )
    d_ent = load_drake_entities(drake)
    l_ent, log_stats = load_log_entities(log)

    tags = log_stats["pool_tags"]
    core_pool = [s for s, t in zip(l_ent, tags) if t == "core"]
    core_xcel = [s for s, t in zip(l_ent, tags) if t in ("core", "xcel_blank_first")]
    all_pool = l_ent

    m_core, res_core, _ = match_entities(d_ent, core_pool)
    m_core_xcel, res_cx, _ = match_entities(d_ent, core_xcel)
    m_all, res_all, _ = match_entities(d_ent, all_pool)

    # How many of core+xcel matches are only possible via xcel blank first?
    matched_core_ids = {m[0].id for m in m_core}
    matched_cx_ids = {m[0].id for m in m_core_xcel}
    only_xcel = len(matched_cx_ids - matched_core_ids)

    # Person-side regression check
    wb = openpyxl.load_workbook(drake, read_only=True, data_only=True)
    ws = wb.active
    drake_names = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i == 1:
            continue
        vals = list(row)
        name = str(vals[1] or "").strip() if len(vals) > 1 else ""
        if not name or name.upper().startswith("TOTAL"):
            continue
        rtype = str(vals[2] or "").strip().upper() if len(vals) > 2 else ""
        if rtype in config.ENTITY_DRAKE_TYPES:
            continue
        drake_names.append(name)
    wb.close()
    wb = openpyxl.load_workbook(log, read_only=True, data_only=True)
    ws = wb[config.SHEET_INDIVIDUALS]
    log_rows = []
    for i, row in enumerate(ws.iter_rows(values_only=True), 1):
        if i < config.LOG_DATA_START_ROW:
            continue
        vals = list(row)
        last = str(vals[config.LOG_LAST_COL] or "").strip() if len(vals) > config.LOG_LAST_COL else ""
        first = (
            str(vals[config.LOG_FIRST_COL] or "").strip()
            if len(vals) > config.LOG_FIRST_COL
            else ""
        )
        if last and first:
            log_rows.append((last, first))
    wb.close()
    person = drake_log_match_rate(drake_names, log_rows)

    before_rate = 0.6087
    primary = m_core_xcel  # recommended pool
    after_rate = len(primary) / len(d_ent) if d_ent else 0.0
    residual_causes = characterize_residual(res_cx, core_xcel)

    out = {
        "before_rate": before_rate,
        "before_matched": 70,
        "before_total": 115,
        "after_pool": "core_business_lists + xcel_blank_first",
        "after_matched": len(primary),
        "after_total": len(d_ent),
        "after_rate": round(after_rate, 4),
        "core_only_matched": len(m_core),
        "core_only_rate": round(len(m_core) / len(d_ent), 4),
        "all_including_ext_matched": len(m_all),
        "all_including_ext_rate": round(len(m_all) / len(d_ent), 4),
        "matched_only_via_xcel_blank_first": only_xcel,
        "xcel_blank_first_count": log_stats["xcel_blank_first"],
        "core_rows": log_stats["core_rows"],
        "ext_rows": log_stats["ext_rows"],
        "residual_count": len(res_cx),
        "residual_causes": residual_causes,
        "person_rate": person["rate"],
        "person_matched": person["matched"],
        "person_total": person["drake"],
        "person_unchanged": round(person["rate"], 4) == 0.9154,
        "rule_justification": {
            "ampersand_as_name": {
                "examples_in_data": ">=2 (J&H, E&G, J & J on business sheets)",
                "count_drake_amp": 6,
            },
            "trailing_INC_LLC": {"examples_in_data": "INC>=57 Drake, LLC>=33 Drake"},
            "CORPORATION": {"examples_in_data": "Drake=1, TaxLog>=6"},
            "leading_THE": {
                "examples_in_data": "TaxLog THE PRINCESS MATTRESS family >=2"
            },
            "leading_zero_variant": {
                "examples_in_data": "011 INTERNATIONAL on Drake and 1120 CORP LIST"
            },
            "hyphen_digits": {
                "examples_in_data": "digit-hyphen patterns on business sheets (>=2)"
            },
            "DBA": {
                "examples_in_data": 0,
                "note": "Implemented per spec but inert — 0 occurrences in sources",
            },
            "L.L.C_dot_form": {
                "note": "Punctuation collapse into LLC suffix (LLC>=33 justifies)"
            },
            "exclude_EXT_from_primary_pool": {
                "examples_in_data": "EXT sheets add collision noise; core lists are the firm roster",
                "note": "EXT retained in secondary measurement only",
            },
        },
    }
    Path(r"T:\audit\output\m92_entity_rate.json").write_text(
        json.dumps(out, indent=2), encoding="utf-8"
    )
    print(json.dumps(out, indent=2))


if __name__ == "__main__":
    main()
