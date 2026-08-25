"""M9.5 — regenerate findings workbook with M9 summary sections."""

from __future__ import annotations

import json
from collections import defaultdict
from pathlib import Path

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from audit.db import connect_audit
from audit.util import stamp_day

HEADER_FILL = PatternFill("solid", fgColor="1E293B")
HEADER_FONT = Font(name="Calibri", bold=True, color="FFFFFF", size=11)
DATA_FONT = Font(name="Calibri", size=11)
BOLD = Font(name="Calibri", bold=True, size=11)
CENTER = Alignment(horizontal="center", vertical="center")
THIN = Border(bottom=Side(style="thin", color="E2E8F0"))
PASTEL = {
    "NEEDS_HUMAN": "FECACA",
    "DUPLICATE_CLIENT": "FDE68A",
    "MISSING_IN_TAXOPS": "FECACA",
    "PHANTOM_IN_TAXOPS": "E0E7FF",
    "PREPARED_NOT_LOGGED": "FEF9C3",
    "LOGGED_NOT_PREPARED": "CCFBF1",
    "SPOUSE_MISSING": "FCE7F3",
    "SPOUSE_STORE_DIVERGENCE": "E0F2FE",
    "NAME_TRUNCATED": "F1F5F9",
    "CONFIRMED_MATCH": "D1FAE5",
    "FIELD_MISMATCH": "FEF3C7",
}


def _hdr(ws, cols):
    for i, (label, width) in enumerate(cols, 1):
        cell = ws.cell(1, i, label)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 22


def _row(ws, r, vals, fill=None):
    for ci, v in enumerate(vals, 1):
        cell = ws.cell(r, ci, v)
        cell.font = DATA_FONT
        cell.border = THIN
        if fill and ci == 1:
            cell.fill = PatternFill("solid", fgColor=fill)


def write_m9_workbook(
    audit_db: Path,
    run_id: int,
    out_dir: Path,
    *,
    m91: dict,
    m92: dict,
    m93: dict,
    m94: dict,
) -> Path:
    out_dir.mkdir(parents=True, exist_ok=True)
    # New timestamped name — do not overwrite prior
    from datetime import datetime, timezone

    stamp = datetime.now(timezone.utc).strftime("%Y%m%dT%H%M%SZ")
    out = out_dir / f"client_audit_findings_M9_{stamp}_run{run_id}.xlsx"

    conn = connect_audit(audit_db)
    try:
        findings = conn.execute(
            "SELECT * FROM audit_finding WHERE run_id=? ORDER BY severity ASC, id ASC",
            (run_id,),
        ).fetchall()
        by_type: dict[str, list] = defaultdict(list)
        for f in findings:
            by_type[f["finding_type"]].append(f)

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Summary"
        _hdr(ws, [("Section", 28), ("Key", 72), ("Count", 14)])
        row = 2

        def section(title):
            nonlocal row
            ws.cell(row, 1, title).font = BOLD
            row += 1

        def kv(sec, key, count, fill=None):
            nonlocal row
            _row(ws, row, [sec, key, count], fill)
            row += 1

        section("M9.1 CSM export diff")
        kv("M9.1", f"older_count={m91.get('older_count')}", m91.get("older_count"))
        kv("M9.1", f"newer_count={m91.get('newer_count')}", m91.get("newer_count"))
        kv("M9.1", "removals", m91.get("removals_count"))
        kv("M9.1", "additions", m91.get("additions_count"))
        rec = m91.get("recommendation") or {}
        kv("M9.1", f"recommendation={rec.get('action')}", "")
        kv("M9.1", f"keep_drift_flag={rec.get('keep_acknowledge_baseline_drift')}", "")
        for tag, cnt in (m91.get("presence_patterns") or {}).items():
            kv("M9.1 presence", tag, cnt)

        section("M9.2 Entity match rate")
        kv("M9.2", "before_rate", m92.get("before_rate"))
        kv("M9.2", "after_rate_core_lists", m92.get("core_only_rate"))
        kv("M9.2", "after_matched_core", m92.get("core_only_matched"))
        kv("M9.2", "after_rate_core_plus_xcel_blank_first", m92.get("after_rate"))
        kv("M9.2", "xcel_blank_first_rows", m92.get("xcel_blank_first_count"))
        kv("M9.2", "residual", m92.get("residual_count"))
        for cause, cnt in (m92.get("residual_causes") or {}).items():
            kv("M9.2 residual", cause, cnt)
        kv("M9.2", "person_rate_unchanged", m92.get("person_rate"))

        section("M9.3 TY2025 unmatched causal cross-tab")
        kv("M9.3", "unmatched_total", m93.get("unmatched_total"))
        kv("M9.3", "top_coverage", m93.get("top_coverage"))
        for m, c in (m93.get("month_histogram") or {}).items():
            kv("M9.3 month", m, c)
        for t, c in (m93.get("bulk_timestamps_gt_20") or {}).items():
            kv("M9.3 bulk_ts", t, c)
        for p in m93.get("populations") or []:
            kv("M9.3 population", f"{p['hypothesis'][:80]}", p["count"])

        section("M9.4 Spouse store divergence split")
        for cls, cnt in (m94.get("classes") or {}).items():
            kv("M9.4", cls, cnt)
        kv("M9.4", "CONFLICT", m94.get("conflict_count"))
        kv("M9.4", "remediation", m94.get("remediation"))
        kv("M9.4", m94.get("plain_statement", ""), "")

        section("Findings by type")
        for ftype, rows_ in sorted(by_type.items(), key=lambda x: x[1][0]["severity"]):
            kv("FINDING", ftype, len(rows_), PASTEL.get(ftype))

        # Finding tabs
        for ftype, rows_ in sorted(by_type.items(), key=lambda x: x[1][0]["severity"]):
            wsf = wb.create_sheet(ftype[:31])
            _hdr(
                wsf,
                [
                    ("ID", 8),
                    ("Subtype", 28),
                    ("Severity", 10),
                    ("Subject kind", 16),
                    ("Subject id", 12),
                    ("Detail JSON", 48),
                    ("Source refs", 36),
                ],
            )
            fill = PASTEL.get(ftype)
            for ri, f in enumerate(rows_, 2):
                vals = [
                    f["id"],
                    f["subtype"],
                    f["severity"],
                    f["subject_kind"],
                    f["subject_id"],
                    f["detail_json"],
                    f["source_refs"],
                ]
                for ci, v in enumerate(vals, 1):
                    cell = wsf.cell(ri, ci, v)
                    cell.font = DATA_FONT
                    cell.border = THIN
                    if fill and ci == 2:
                        cell.fill = PatternFill("solid", fgColor=fill)

        # M9.1 detail tab (names — gitignored workbook only)
        wsd = wb.create_sheet("M91_diff_rows")
        _hdr(
            wsd,
            [
                ("Direction", 18),
                ("Type", 8),
                ("Status", 14),
                ("Preparer", 18),
                ("Started", 14),
                ("Completed", 14),
                ("Last Change", 22),
                ("Changed By", 14),
                ("Total Bill", 12),
                ("in_log", 8),
                ("in_taxops", 10),
                ("in_ty2025", 10),
                ("Client Name", 40),
                ("ID Last4", 12),
            ],
        )
        ri = 2
        for item in (m91.get("removals") or []) + (m91.get("additions") or []):
            r = item["row"]
            p = item["presence"]
            vals = [
                item["direction"],
                r.get("Type"),
                r.get("Status"),
                r.get("Preparer"),
                str(r.get("Started")),
                str(r.get("Completed")),
                str(r.get("Last Change")),
                r.get("Changed By"),
                r.get("Total Bill"),
                p.get("in_tax_log"),
                p.get("in_taxops_client"),
                p.get("in_taxops_ty2025"),
                r.get("Client Name"),
                r.get("_id_last4"),
            ]
            for ci, v in enumerate(vals, 1):
                cell = wsd.cell(ri, ci, v)
                cell.font = DATA_FONT
                cell.border = THIN
            ri += 1

        wb.save(out)
        return out
    finally:
        conn.close()


if __name__ == "__main__":
    out_dir = Path(r"T:\audit\output")
    m91 = json.loads((out_dir / "m91_export_diff.json").read_text(encoding="utf-8"))
    # rebuild console-safe recommendation block onto full diff
    from audit.export_diff import console_summary

    m91_summary = console_summary(m91)
    m92 = json.loads((out_dir / "m92_entity_rate.json").read_text(encoding="utf-8"))
    m93 = json.loads((out_dir / "m93_gap_causal.json").read_text(encoding="utf-8"))
    m94 = json.loads((out_dir / "m94_store_split.json").read_text(encoding="utf-8"))

    dbs = sorted(Path(r"T:\audit").glob("audit_*.sqlite"), key=lambda p: p.stat().st_mtime)
    db = dbs[-1]
    from audit.db import connect_audit as ca

    conn = ca(db)
    run_id = int(conn.execute("SELECT MAX(id) FROM audit_run").fetchone()[0])
    conn.close()

    path = write_m9_workbook(
        db, run_id, out_dir, m91=m91_summary | {"removals": m91["removals"], "additions": m91["additions"]}, m92=m92, m93=m93, m94=m94
    )
    print(json.dumps({"workbook": str(path), "run_id": run_id}, indent=2))
