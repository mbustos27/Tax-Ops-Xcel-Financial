import json, csv, io, re
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

OUT = Path(r"T:\audit\investigation")

schema = json.loads((OUT/"I0-schema-dump.json").read_text(encoding="utf-8"))
live = next(x for x in schema if Path(x["path"]).name == "taxops.db")
identity_keywords = re.compile(
    r"client|return|spouse|dependent|ssn|ein|phone|email|address|log_number|"
    r"drake|prefill|csm|taxpayer|billing|name|person|household|import|staging|"
    r"match|link|audit",
    re.I,
)
identity_tables = []
for obj in live["sqlite_master"]:
    if obj["type"] != "table":
        continue
    name = obj["name"]
    sql = obj["sql"] or ""
    if identity_keywords.search(name) or identity_keywords.search(sql):
        identity_tables.append({"name": name, "row_count": live["row_counts"].get(name), "sql": sql})
(OUT/"I0-identity-tables.json").write_text(json.dumps(identity_tables, indent=2), encoding="utf-8")
print("identity tables", len(identity_tables))
for t in identity_tables:
    print(f"  {t['row_count']:>6}  {t['name']}")

files = [
    Path(r"T:\taxops\CSVFILES\2024 CLIENTS.xlsx"),
    Path(r"T:\taxops\CSVFILES\TY2024S.csv"),
    Path(r"T:\taxops\CSVFILES\TY2024S2439-5405.csv"),
    Path(r"T:\taxops\CSVFILES\TY2024S8867-end.csv"),
    Path(r"T:\taxops\CSVFILES\TY2024Spouses.csv"),
    Path(r"T:\taxops\CSVFILES\TAXPAYERspouse25.csv"),
    Path(r"C:\TaxOps\taxops\CSVFILES\TY2024Spouses.csv"),
    Path(r"C:\TaxOps\taxops\CSVFILES\TAXPAYERspouse25.csv"),
    Path(r"C:\Users\Windows 10\Desktop\CLIENTS.xlsx"),
    Path(r"C:\Users\Windows 10\OneDrive - Xcel Financial Services LLC\Shared\Logs\TAX LOG 2025 Live.xlsx"),
]
od = Path(r"C:\Users\Windows 10\OneDrive - Xcel Financial Services LLC")
if od.exists():
    for p in od.rglob("CLIENTS.xlsx"):
        files.append(p)
        print("found", p)

for p in Path(r"T:\audit").rglob("*"):
    if p.suffix.lower() in {".csv",".xlsx",".xls",".tsv"} and "investigation" not in str(p):
        files.append(p)

seen=set(); uniq=[]
for p in files:
    k=str(p)
    if k not in seen:
        seen.add(k); uniq.append(p)

meta_list=[]
for p in uniq:
    if not p.exists():
        meta_list.append({"path":str(p),"exists":False})
        print("MISSING", p)
        continue
    st=p.stat()
    meta={"path":str(p),"size":st.st_size,"mtime":datetime.fromtimestamp(st.st_mtime).isoformat(sep=" ",timespec="seconds"),"ext":p.suffix.lower()}
    try:
        if p.suffix.lower()==".csv":
            text=None
            for enc in ("utf-8-sig","utf-8","cp1252","latin-1"):
                try:
                    text=p.read_text(encoding=enc); meta["encoding"]=enc; break
                except Exception:
                    pass
            first=text.splitlines()[0] if text and text.splitlines() else ""
            meta["header_raw"]=first
            delim="," if first.count(",")>=first.count("\t") else "\t"
            meta["delimiter"]=repr(delim)
            rows=list(csv.reader(io.StringIO(text), delimiter=delim))
            meta["row_count_incl_header"]=len(rows)
            meta["data_row_count"]=max(0,len(rows)-1)
            meta["header_cols"]=rows[0] if rows else []
            meta["col_count"]=len(rows[0]) if rows else 0
        else:
            wb=load_workbook(p, read_only=True, data_only=True)
            sheets=[]
            for sn in wb.sheetnames:
                ws=wb[sn]
                it=ws.iter_rows(values_only=True)
                header=next(it, None)
                n=sum(1 for _ in it)
                sheets.append({
                    "sheet":sn,
                    "header":[("" if c is None else str(c)) for c in (header or [])],
                    "data_row_count":n,
                    "col_count":len(header) if header else 0,
                })
            wb.close()
            meta["sheets"]=sheets
    except Exception as e:
        meta["error"]=str(e)
    meta_list.append(meta)
    print("OK", p.name)

(OUT/"I0-ingestion-files.json").write_text(json.dumps(meta_list, indent=2), encoding="utf-8")
print("Wrote", len(meta_list), "files")
