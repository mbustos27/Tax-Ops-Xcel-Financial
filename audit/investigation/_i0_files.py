import json, csv, re
from pathlib import Path
from datetime import datetime
from openpyxl import load_workbook

OUT = Path(r"T:\audit\investigation")
schema = json.loads((OUT/"I0-schema-dump.json").read_text(encoding="utf-8"))

# Identity-ish table heuristic from live taxops.db snapshot
live = next(x for x in schema if x["path"].endswith("taxops.db"))
identity_keywords = re.compile(
    r"client|return|spouse|dependent|ssn|ein|phone|email|address|log_number|"
    r"drake|prefill|csm|taxpayer|billing|name|person|household|import|staging|"
    r"match|link|audit",
    re.I,
)
identity_tables = []
for obj in live["sqlite_master"]:
    if obj["type"] != "table":
        continue
    name = obj["name"]
    sql = obj["sql"] or ""
    if identity_keywords.search(name) or identity_keywords.search(sql):
        identity_tables.append({
            "name": name,
            "row_count": live["row_counts"].get(name),
            "sql": sql,
        })

(OUT/"I0-identity-tables.json").write_text(json.dumps(identity_tables, indent=2), encoding="utf-8")
print("identity tables:", len(identity_tables))
for t in identity_tables:
    print(f"  {t['row_count']:>6}  {t['name']}")

# Ingestion files inventory
candidates = []
roots = [
    Path(r"T:\taxops\CSVFILES"),
    Path(r"C:\TaxOps\taxops\CSVFILES"),
    Path(r"T:\audit"),
    Path(r"T:\taxops"),
]
patterns = ["*.csv","*.xlsx","*.xls","*.tsv"]
seen = set()
for root in roots:
    if not root.exists():
        continue
    for pat in patterns:
        for p in root.rglob(pat):
            # skip node_modules / venv / huge unrelated
            s = str(p).lower()
            if any(x in s for x in [".venv","node_modules","__pycache__",".git"]):
                continue
            if str(p) in seen:
                continue
            seen.add(str(p))
            candidates.append(p)

# Prefer likely tax/client sources; still record all under CSVFILES and audit
def interesting(p: Path) -> bool:
    s = str(p).lower()
    name = p.name.lower()
    if "csvfiles" in s or r"\audit\\" in s or "/audit/" in s.replace("\\","/"):
        return True
    keys = ["client","ty202","tax","csm","purple","spouse","taxpayer","log","drake","import"]
    return any(k in name for k in keys)

files_meta = []
for p in sorted(candidates, key=lambda x: str(x).lower()):
    if not interesting(p):
        continue
    st = p.stat()
    meta = {
        "path": str(p),
        "size": st.st_size,
        "mtime": datetime.fromtimestamp(st.st_mtime).isoformat(sep=" ", timespec="seconds"),
        "ext": p.suffix.lower(),
    }
    try:
        if p.suffix.lower() == ".csv":
            # try encodings
            text = None
            for enc in ("utf-8-sig","utf-8","cp1252","latin-1"):
                try:
                    text = p.read_text(encoding=enc)
                    meta["encoding"] = enc
                    break
                except Exception:
                    continue
            if text is None:
                meta["error"] = "decode failed"
            else:
                # detect delimiter from first line
                first = text.splitlines()[0] if text.splitlines() else ""
                meta["header_raw"] = first
                # row count via csv
                import io
                # sniff
                sample = "\n".join(text.splitlines()[:5])
                try:
                    dialect = csv.Sniffer().sniff(sample, delimiters=",\t|;")
                    delim = dialect.delimiter
                except Exception:
                    delim = "," if first.count(",") >= first.count("\t") else "\t"
                meta["delimiter"] = repr(delim)
                reader = csv.reader(io.StringIO(text), delimiter=delim)
                rows = list(reader)
                meta["row_count_incl_header"] = len(rows)
                meta["data_row_count"] = max(0, len(rows)-1)
                meta["header_cols"] = rows[0] if rows else []
                meta["col_count"] = len(rows[0]) if rows else 0
        elif p.suffix.lower() in (".xlsx",".xls"):
            wb = load_workbook(p, read_only=True, data_only=True)
            sheets = []
            for sn in wb.sheetnames:
                ws = wb[sn]
                rows = ws.iter_rows(values_only=True)
                header = next(rows, None)
                n = 0
                for _ in rows:
                    n += 1
                sheets.append({
                    "sheet": sn,
                    "header": [str(c) if c is not None else "" for c in (header or [])],
                    "data_row_count": n,
                    "col_count": len(header) if header else 0,
                })
            wb.close()
            meta["sheets"] = sheets
    except Exception as e:
        meta["error"] = str(e)
    files_meta.append(meta)
    print(f"FILE {p.name} rows={meta.get('data_row_count') or meta.get('sheets')}")

(OUT/"I0-ingestion-files.json").write_text(json.dumps(files_meta, indent=2), encoding="utf-8")
print("Wrote I0-ingestion-files.json", len(files_meta))
