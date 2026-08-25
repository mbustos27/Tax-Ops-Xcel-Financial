"""One-shot debug: invoice vs TaxOps vs Tax Log key formats. Read-only."""
from __future__ import annotations

import re
from collections import Counter
from pathlib import Path

import openpyxl

from audit import config
from audit.baseline import load_baseline_memory
from audit.db import connect_taxops_readonly
from audit.invoice_export import DEFAULT_TAXPAYER_INVOICE_PATH, parse_taxpayer_invoice_csv

mem = load_baseline_memory()
inv = parse_taxpayer_invoice_csv(DEFAULT_TAXPAYER_INVOICE_PATH)
l0 = set(inv.l0_ok_invoices)
print("invoice L0 sample:", sorted(l0)[:15], "n=", len(l0))
print("invoice lens:", Counter(len(x) for x in l0))

taxops = Path(mem["authoritative_taxops_path"])
conn = connect_taxops_readonly(taxops)
t_raw = [
    str(r[0]).strip()
    for r in conn.execute(
        "SELECT log_number FROM returns WHERE tax_year=2025 "
        "AND log_number IS NOT NULL AND TRIM(log_number)!=''"
    )
]
print("taxops sample raw:", t_raw[:20])
print("taxops lens:", Counter(len(x) for x in t_raw))
print("taxops startswith 25:", sum(1 for x in t_raw if x.startswith("25")))
tkeys = set(t_raw)
print("overlap invoice∩taxops exact:", len(l0 & tkeys))

# try zero-pad / strip season
def variants(s: str) -> set[str]:
    out = {s, s.lstrip("0") or "0"}
    digits = re.sub(r"\D", "", s)
    if digits:
        out.add(digits)
        out.add(digits.zfill(6))
        if len(digits) == 6 and digits.startswith("25"):
            out.add(digits[2:])  # drop season
        if len(digits) <= 4:
            out.add("25" + digits.zfill(4))
    return out

# best-effort overlap via variants
hit = 0
for inv_no in list(l0)[:]:
    vs = variants(inv_no)
    if vs & tkeys:
        hit += 1
print("overlap via variants (invoice→taxops):", hit)

# reverse: how many taxops match an invoice via pad
hit2 = 0
for t in tkeys:
    if variants(t) & l0:
        hit2 += 1
print("overlap via variants (taxops→invoice):", hit2)

log_path = Path(mem["tax_log_path"])
wb = openpyxl.load_workbook(log_path, read_only=True, data_only=True)
ws = wb[config.SHEET_INDIVIDUALS]
print("--- Tax Log header rows ---")
for i, row in enumerate(ws.iter_rows(values_only=True, max_col=12), 1):
    if i <= 5:
        print(i, list(row))
    else:
        break
# Sample data rows
print("--- data samples ---")
n = 0
for i, row in enumerate(ws.iter_rows(values_only=True, max_col=12), 1):
    if i < 6:
        continue
    vals = list(row)
    last = str(vals[2] or "").strip() if len(vals) > 2 else ""
    first = str(vals[3] or "").strip() if len(vals) > 3 else ""
    if not (last or first):
        continue
    print(
        "row",
        i,
        "A=",
        repr(vals[0])[:20],
        "B=",
        repr(vals[1])[:30],
        "C=",
        last[:25],
        "D=",
        first[:20],
        "E=",
        repr(vals[4])[:15] if len(vals) > 4 else None,
    )
    n += 1
    if n >= 8:
        break

# Find which column looks like 25xxxx
col_hits = Counter()
for i, row in enumerate(ws.iter_rows(values_only=True, max_col=20), 1):
    if i < 6:
        continue
    vals = list(row)
    for ci, v in enumerate(vals):
        s = str(v or "").strip()
        digits = re.sub(r"\D", "", s)
        if re.fullmatch(r"25\d{4}", digits):
            col_hits[ci] += 1
print("cols with 25xxxx pattern counts:", dict(col_hits))
wb.close()
conn.close()
